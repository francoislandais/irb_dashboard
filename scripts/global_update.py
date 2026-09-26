"""Prototype de point d'entrée unique pour planifier les exports Agora Explorer.

Le mode ``preview`` écrit les requêtes SQL sans les exécuter. Le mode ``test``
redirige le SQL vers les tables de test documentées et construit les datasets
depuis le jeu factice local, sans connexion à Hive.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import json
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

if __package__:
    from .export_all_standalone_apps import export_standalone_app
    from .hive_to_dataset import (
        _expand_template_expressions,
        build_hive_query,
        build_kri_hive_query,
        find_kris_using_templates,
    )
else:
    from export_all_standalone_apps import export_standalone_app
    from hive_to_dataset import (
        _expand_template_expressions,
        build_hive_query,
        build_kri_hive_query,
        find_kris_using_templates,
    )


PROJECT_DIRECTORY = Path(__file__).resolve().parents[1]
EXAMPLE_WORKBOOK = PROJECT_DIRECTORY / "outputs" / "global-update-prototype" / "global_update_examples.xlsx"
TEST_ENTITIES_PATH = PROJECT_DIRECTORY / "scripts" / "fixtures" / "global_update_test_entities.json"
DEFAULT_OUTPUT_DIRECTORY = PROJECT_DIRECTORY / "outputs" / "global-update-prototype" / "generated"
LEI_TOKEN = "__GLOBAL_UPDATE_LEI_FILTER__"
FREQUENCIES = {"MONTHLY", "QUARTERLY", "SEMI_ANNUAL", "ANNUAL"}
PERIODS_PER_YEAR = {"MONTHLY": 12, "QUARTERLY": 4, "SEMI_ANNUAL": 2, "ANNUAL": 1}
CONSOLIDATION_MODES = {"HIGHEST", "CONSO", "SOLO", "CONSO+SOLO", "ALL"}
TEST_TABLES = {
    "ITS": "crp_agora.agora_its_bft_current_test",
    "KRI": "crp_agora.agora_dm_imas_kris_raw_test",
}
PRODUCTION_TABLES = {
    "ITS": "crp_agora.agora_its_bft_current",
    "KRI": "crp_agora.agora_dm_imas_kris_raw",
}
OUTPUT_COLUMNS = [
    "table_id",
    "reporting_unit_id",
    "x_axis_rc_code",
    "y_axis_rc_code",
    "z_axis_rc_code",
]


@dataclass(frozen=True)
class Extraction:
    module: str
    selector: str
    history_years: int
    frequency: str
    reference_dates: tuple[str, ...]
    kri_data_point_ids: tuple[str, ...] = ()


@dataclass(frozen=True)
class Application:
    name: str
    leis: tuple[str, ...]
    consolidation: str
    extractions: tuple[Extraction, ...]


def _split_list(value: object) -> list[str]:
    return list(dict.fromkeys(part.strip() for part in str(value or "").split(",") if part.strip()))


def _month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def _subtract_months(value: date, months: int) -> date:
    month_index = value.year * 12 + value.month - 1 - months
    year, month_zero_based = divmod(month_index, 12)
    month = month_zero_based + 1
    return _month_end(year, month)


def _latest_period_end(frequency: str, as_of: date) -> date:
    candidates: list[date] = []
    if frequency == "MONTHLY":
        candidates = [_month_end(as_of.year, as_of.month)]
        if candidates[0] > as_of:
            candidates = [_subtract_months(candidates[0], 1)]
    elif frequency == "QUARTERLY":
        candidates = [date(as_of.year, month, calendar.monthrange(as_of.year, month)[1]) for month in (3, 6, 9, 12)]
    elif frequency == "SEMI_ANNUAL":
        candidates = [date(as_of.year, 6, 30), date(as_of.year, 12, 31)]
    elif frequency == "ANNUAL":
        candidates = [date(as_of.year, 12, 31)]
    else:
        raise ValueError(f"Fréquence inconnue : {frequency}")
    available = [candidate for candidate in candidates if candidate <= as_of]
    if available:
        return available[-1]
    return date(as_of.year - 1, 12, 31)


def _reference_dates(frequency: str, history_periods: int, as_of: date) -> tuple[str, ...]:
    latest = _latest_period_end(frequency, as_of)
    interval_months = {"MONTHLY": 1, "QUARTERLY": 3, "SEMI_ANNUAL": 6, "ANNUAL": 12}[frequency]
    dates = [_subtract_months(latest, interval_months * offset) for offset in range(history_periods - 1, -1, -1)]
    return tuple(item.isoformat() for item in dates)


def _read_configuration(path: Path, as_of: date) -> tuple[Application, ...]:
    if not path.is_file():
        raise ValueError(f"Classeur de paramétrage introuvable : {path}")
    workbook = load_workbook(path, read_only=False, data_only=True)
    applications: list[Application] = []
    lei_pattern = re.compile(r"^[A-Z0-9]{20}$")
    for sheet in workbook.worksheets:
        leis = _split_list(sheet["B1"].value)
        consolidation = str(sheet["B2"].value or "").strip().upper()
        if not leis or any(not lei_pattern.fullmatch(lei.upper()) for lei in leis):
            raise ValueError(f"{sheet.title}: B1 doit contenir une liste de LEI de 20 caractères, séparés par des virgules.")
        leis = [lei.upper() for lei in leis]
        if consolidation not in CONSOLIDATION_MODES:
            raise ValueError(f"{sheet.title}: B2 doit être l'un de {', '.join(sorted(CONSOLIDATION_MODES))}.")
        expected_headers = ["module", "template selector", "history years", "frequency"]
        actual_headers = [str(sheet.cell(4, col).value or "").strip().lower() for col in range(1, 5)]
        if actual_headers != expected_headers:
            raise ValueError(f"{sheet.title}: les en-têtes A4:D4 doivent être {expected_headers}.")
        extractions: list[Extraction] = []
        for row_number in range(5, sheet.max_row + 1):
            module, selector, periods, frequency = [sheet.cell(row_number, col).value for col in range(1, 5)]
            if all(value is None or str(value).strip() == "" for value in (module, selector, periods, frequency)):
                continue
            module = str(module or "").strip().upper()
            selector = str(selector or "").strip()
            frequency = str(frequency or "").strip().upper().replace("-", "_").replace(" ", "_")
            try:
                history_years = int(periods)
            except (TypeError, ValueError) as error:
                raise ValueError(f"{sheet.title}, ligne {row_number}: History years doit être un entier.") from error
            if not module or not selector:
                raise ValueError(f"{sheet.title}, ligne {row_number}: Module et Template selector sont requis.")
            if frequency not in FREQUENCIES:
                raise ValueError(f"{sheet.title}, ligne {row_number}: fréquence invalide ({frequency}).")
            if history_years < 1:
                raise ValueError(f"{sheet.title}, ligne {row_number}: History years doit être supérieur ou égal à 1.")
            kri_ids: tuple[str, ...] = ()
            if module == "KRI":
                kri_ids = tuple(find_kris_using_templates([selector]))
                if not kri_ids:
                    raise ValueError(f"{sheet.title}, ligne {row_number}: aucun KRI ne dépend de {selector!r}.")
            history_periods = history_years * PERIODS_PER_YEAR[frequency]
            dates = _reference_dates(frequency, history_periods, as_of)
            extractions.append(Extraction(module, selector, history_years, frequency, dates, kri_ids))
        if not extractions:
            raise ValueError(f"{sheet.title}: au moins une ligne d'extraction est requise.")
        applications.append(Application(sheet.title, tuple(leis), consolidation, tuple(extractions)))
    workbook.close()
    if not applications:
        raise ValueError("Le classeur ne contient aucun onglet d'application.")
    return tuple(applications)


def _sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _selected_levels(mode: str, highest_level: str) -> tuple[str, ...]:
    if mode == "HIGHEST":
        return (highest_level,)
    if mode == "CONSO+SOLO":
        return ("CONSO", "SOLO")
    if mode == "ALL":
        return ("CONSO", "SOLO", "LIQUIDITY_SUBGROUP")
    return (mode,)


def _apply_consolidation_filter(sql: str, mode: str) -> str:
    highest_filter = "AND is_highest_cons = 'Y'"
    if mode == "HIGHEST":
        return sql
    if highest_filter not in sql:
        raise ValueError("Le générateur Hive existant ne contient pas le filtre is_highest_cons attendu.")
    if mode == "ALL":
        return sql.replace(highest_filter, "")
    levels = _selected_levels(mode, "")
    level_filter = "AND consolidation_level IN (" + ", ".join(_sql_literal(item) for item in levels) + ")"
    return sql.replace(highest_filter, level_filter)


def _replace_lei_filter(sql: str, leis: Iterable[str]) -> str:
    pattern = re.compile(r"jst_code\s+IN\s*\(\s*'" + re.escape(LEI_TOKEN) + r"'\s*\)", re.IGNORECASE)
    lei_filter = "lei IN (" + ", ".join(_sql_literal(item) for item in leis) + ")"
    updated, count = pattern.subn(lei_filter, sql, count=1)
    if count != 1:
        raise ValueError("Impossible de remplacer le filtre JST du générateur existant par le filtre LEI.")
    return updated


def _build_extraction_sql(extraction: Extraction, application: Application, source_table: str) -> str:
    if extraction.module == "KRI":
        sql = build_kri_hive_query(extraction.kri_data_point_ids, extraction.reference_dates, [LEI_TOKEN])
        sql = sql.replace(PRODUCTION_TABLES["KRI"], source_table)
        sql = sql.replace("jst_code AS reporting_unit_id", "CONCAT(lei, '_', consolidation_level) AS reporting_unit_id")
        sql = sql.replace("GROUP BY\n    jst_code,\n    kri_data_point_id", "GROUP BY\n    lei,\n    consolidation_level,\n    jst_code,\n    kri_data_point_id")
        sql = sql.replace("ORDER BY\n    jst_code,\n    kri_data_point_id", "ORDER BY\n    lei,\n    consolidation_level,\n    jst_code,\n    kri_data_point_id")
    else:
        sql = build_hive_query([extraction.selector], extraction.reference_dates, [LEI_TOKEN], module_id=extraction.module)
        sql = sql.replace(PRODUCTION_TABLES["ITS"], source_table)
        sql = sql.replace("jst_code AS reporting_unit_id", "CONCAT(lei, '_', consolidation_level) AS reporting_unit_id")
        sql = sql.replace("AS table_id,\n        jst_code,", "AS table_id,\n        lei,\n        consolidation_level,\n        jst_code,")
        sql = sql.replace("GROUP BY\n    table_id,\n    jst_code,", "GROUP BY\n    table_id,\n    lei,\n    consolidation_level,\n    jst_code,")
        sql = sql.replace("ORDER BY\n    table_id,\n    jst_code,", "ORDER BY\n    table_id,\n    lei,\n    consolidation_level,\n    jst_code,")
    sql = _replace_lei_filter(sql, application.leis)
    return _apply_consolidation_filter(sql, application.consolidation)


def _build_institution_sql(application: Application, table: str) -> str:
    sql = f"""SELECT DISTINCT
    CONCAT(lei, '_', consolidation_level) AS institution_id,
    lei,
    jst_code,
    institution_name,
    consolidation_level
FROM {table}
WHERE lei IN ({', '.join(_sql_literal(lei) for lei in application.leis)})
  AND is_group_head = 'Y'"""
    if application.consolidation == "HIGHEST":
        sql += "\n  AND is_highest_cons = 'Y'"
    elif application.consolidation != "ALL":
        levels = _selected_levels(application.consolidation, "")
        sql += "\n  AND consolidation_level IN (" + ", ".join(_sql_literal(level) for level in levels) + ")"
    return sql + "\nORDER BY lei, consolidation_level\n"


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")
    return cleaned or "application"


def _write_query_index(output: Path, manifest: dict) -> None:
    lines = ["# Requêtes générées", "", f"Mode : `{manifest['mode']}`", f"Date de calcul : `{manifest['as_of']}`", ""]
    for application in manifest["applications"]:
        lines.extend([f"## {application['name']}", "", f"Niveau : `{application['consolidation']}`", f"LEI : {', '.join(application['leis'])}", ""])
        lines.append(f"- [Requête des métadonnées institutionnelles]({application['metadata_query']})")
        for extraction in application["extractions"]:
            lines.append(
                f"- [Extraction {extraction['index']:02d} — {extraction['module']} / {extraction['selector']}]"
                f"({extraction['query']}) — {extraction['frequency']}, {extraction['history_years']} ans"
                f" ({extraction['history_periods']} périodes)"
                f" ({extraction['reference_dates'][0]} → {extraction['reference_dates'][-1]})"
            )
        if application.get("dataset"):
            lines.extend([f"- Dataset de simulation : `{application['dataset']}`", f"- Dictionnaire : `{application['institution_dictionary']}`", f"- Application : `{application['html_app']}`"])
        lines.append("")
    lines.extend(["Le mode `test` écrit des requêtes dirigées vers les tables suffixées `_test`, mais les résultats sont construits depuis la fixture locale et aucune connexion Hive n'est ouverte.", ""])
    output.write_text("\n".join(lines), encoding="utf-8")


def _read_test_fixture(path: Path) -> dict:
    if not path.is_file():
        raise ValueError(f"Fixture du mode test introuvable : {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _all_month_ends(start_iso: str, end_iso: str) -> tuple[str, ...]:
    start = date.fromisoformat(start_iso)
    end = date.fromisoformat(end_iso)
    result = []
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        candidate = _month_end(year, month)
        if start <= candidate <= end:
            result.append(candidate.isoformat())
        month += 1
        if month == 13:
            month = 1
            year += 1
    return tuple(result)


def _selector_matches(value: str, selector: str) -> bool:
    if selector.endswith("%"):
        return value.startswith(selector[:-1])
    return value == selector


def _selected_fixture_templates(selector: str, available: Iterable[str]) -> list[str]:
    included, excluded = _expand_template_expressions([selector])
    return [
        table_id for table_id in available
        if any(_selector_matches(table_id, pattern) for pattern in included)
        and not any(_selector_matches(table_id, pattern) for pattern in excluded)
    ]


def _dictionary_rows(application: Application, fixture: dict) -> list[dict[str, str]]:
    result = []
    entities_by_lei = {item["lei"]: item for item in fixture["entities"]}
    for lei in application.leis:
        entity = entities_by_lei.get(lei)
        if not entity:
            continue
        for level in _selected_levels(application.consolidation, entity["highest_level"]):
            jst_code = entity["jst_by_level"].get(level)
            if not jst_code:
                continue
            result.append({
                "Institution ID": f"{lei}_{level}",
                "JST code": jst_code,
                "Institution Name": entity["institution_name"],
                "Consolidation Level": level,
            })
    return result


def _build_dummy_dataset(application: Application, fixture: dict) -> tuple[list[str], list[dict[str, str]]]:
    all_dates = sorted({ref for item in application.extractions for ref in item.reference_dates})
    headers = OUTPUT_COLUMNS + ["ref_" + ref.replace("-", "_") for ref in all_dates] + ["extraction_timestamp"]
    dictionary = _dictionary_rows(application, fixture)
    selected_ids = {item["Institution ID"] for item in dictionary}
    entities_by_lei = {item["lei"]: item for item in fixture["entities"]}
    records: dict[tuple[str, ...], dict[str, str]] = {}
    available_dates = _all_month_ends(fixture["period_start"], fixture["period_end"])
    module_templates = fixture["templates"]
    for extraction in application.extractions:
        if extraction.module == "KRI":
            templates = list(extraction.kri_data_point_ids)
            is_kri = True
        else:
            templates = _selected_fixture_templates(extraction.selector, module_templates.get(extraction.module, []))
            is_kri = False
        date_set = set(extraction.reference_dates)
        for lei in application.leis:
            entity = entities_by_lei.get(lei)
            if not entity:
                continue
            for level in _selected_levels(application.consolidation, entity["highest_level"]):
                jst_code = entity["jst_by_level"].get(level)
                institution_id = f"{lei}_{level}"
                if not jst_code or institution_id not in selected_ids:
                    continue
                for selected_template in templates:
                    table_id = "KRI" if is_kri else selected_template
                    for period_index, reference_date in enumerate(available_dates, start=1):
                        if reference_date not in date_set:
                            continue
                        x_code = ""
                        y_code = selected_template if is_kri else "0010"
                        z_code = ""
                        key = (table_id, institution_id, x_code, y_code, z_code)
                        record = records.setdefault(key, {
                            "table_id": table_id,
                            "reporting_unit_id": institution_id,
                            "x_axis_rc_code": x_code,
                            "y_axis_rc_code": y_code,
                            "z_axis_rc_code": z_code,
                            **{"ref_" + item.replace("-", "_"): "" for item in all_dates},
                            "extraction_timestamp": date.today().isoformat(),
                        })
                        template_seed = sum(ord(character) for character in table_id + y_code)
                        entity_seed = sum(ord(character) for character in lei)
                        seed = template_seed + entity_seed + (0 if level == "CONSO" else 50 if level == "SOLO" else 90)
                        record["ref_" + reference_date.replace("-", "_")] = str(seed + period_index)
    populated = [
        column for column in headers
        if not column.startswith("ref_") or any(row.get(column, "") != "" for row in records.values())
    ]
    return populated, list(records.values())


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def global_update(
    config_path: str | Path | None = None,
    *,
    mode: str = "preview",
    output_directory: str | Path | None = None,
    as_of: date | None = None,
) -> dict:
    """Point d'entrée du prototype. ``preview`` génère le SQL; ``test`` simule les exports localement."""

    mode = str(mode).strip().lower()
    if mode not in {"preview", "test"}:
        raise ValueError("mode doit être 'preview' ou 'test'. Le mode Hive réel n'est pas activé dans ce prototype.")
    calculation_date = as_of or date.today()
    workbook_path = Path(config_path) if config_path else EXAMPLE_WORKBOOK
    output_root = Path(output_directory) if output_directory else DEFAULT_OUTPUT_DIRECTORY / mode
    applications = _read_configuration(workbook_path, calculation_date)
    fixture = _read_test_fixture(TEST_ENTITIES_PATH) if mode == "test" else None
    query_root = output_root / "queries"
    manifest: dict = {"mode": mode, "as_of": calculation_date.isoformat(), "applications": []}
    for application in applications:
        folder = _safe_name(application.name)
        app_query_dir = query_root / folder
        app_query_dir.mkdir(parents=True, exist_ok=True)
        query_table = TEST_TABLES if mode == "test" else PRODUCTION_TABLES
        metadata_relative = (Path("queries") / folder / "institution_metadata.sql").as_posix()
        metadata_sql = _build_institution_sql(application, query_table["ITS"])
        (output_root / metadata_relative).parent.mkdir(parents=True, exist_ok=True)
        (output_root / metadata_relative).write_text(metadata_sql, encoding="utf-8")
        application_manifest = {
            "name": application.name,
            "leis": list(application.leis),
            "consolidation": application.consolidation,
            "metadata_query": metadata_relative,
            "extractions": [],
        }
        for index, extraction in enumerate(application.extractions, start=1):
            source_table = query_table["KRI" if extraction.module == "KRI" else "ITS"]
            sql = _build_extraction_sql(extraction, application, source_table)
            query_relative = (Path("queries") / folder / f"extraction_{index:02d}.sql").as_posix()
            (output_root / query_relative).write_text(sql, encoding="utf-8")
            application_manifest["extractions"].append({
                "index": index,
                "module": extraction.module,
                "selector": extraction.selector,
                "history_years": extraction.history_years,
                "history_periods": len(extraction.reference_dates),
                "frequency": extraction.frequency,
                "reference_dates": list(extraction.reference_dates),
                "kri_data_point_ids": list(extraction.kri_data_point_ids),
                "query": query_relative,
                "source_table": source_table,
            })
        if mode == "test":
            app_output_dir = output_root / "applications"
            safe_file_name = folder.lower()
            dataset_path = app_output_dir / "datasets" / f"{safe_file_name}.csv"
            dictionary_path = app_output_dir / "institutions" / f"{safe_file_name}_institution_dictionary.csv"
            fields, rows = _build_dummy_dataset(application, fixture)
            dictionary_rows = _dictionary_rows(application, fixture)
            _write_csv(dataset_path, fields, rows)
            dictionary_fields = ["Institution ID", "JST code", "Institution Name", "Consolidation Level"]
            _write_csv(dictionary_path, dictionary_fields, dictionary_rows)
            html_path = app_output_dir / "html" / f"Agora Explorer_{folder}.html"
            export_standalone_app(
                dataset_path,
                html_path,
                institution_dictionary_file_path=dictionary_path,
                app_name=f"Agora Explorer — {application.name}",
            )
            application_manifest.update({
                "dataset": str(dataset_path.relative_to(output_root)),
                "institution_dictionary": str(dictionary_path.relative_to(output_root)),
                "html_app": str(html_path.relative_to(output_root)),
                "simulated_rows": len(rows),
            })
        manifest["applications"].append(application_manifest)
    output_root.mkdir(parents=True, exist_ok=True)
    manifest_path = output_root / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _write_query_index(output_root / "query_index.md", manifest)
    return {"output_directory": str(output_root), "manifest_path": str(manifest_path), "mode": mode, "applications": len(applications)}


def _main() -> None:
    parser = argparse.ArgumentParser(description="Génère les requêtes d'export Agora Explorer ou simule les exports.")
    parser.add_argument("--config", type=Path, default=None, help="Classeur XLSX de paramétrage")
    parser.add_argument("--mode", choices=("preview", "test"), default="preview")
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--as-of", type=date.fromisoformat, default=None, help="Date de calcul YYYY-MM-DD (utile aux essais reproductibles)")
    args = parser.parse_args()
    try:
        result = global_update(args.config, mode=args.mode, output_directory=args.output, as_of=args.as_of)
    except Exception as error:
        parser.exit(2, f"global_update: {error}\n")
    print(f"{result['applications']} application(s) planifiée(s) en mode {result['mode']}")
    print(f"Plan et requêtes : {result['output_directory']}")
    print(f"Index : {result['output_directory']}/query_index.md")


if __name__ == "__main__":
    _main()
