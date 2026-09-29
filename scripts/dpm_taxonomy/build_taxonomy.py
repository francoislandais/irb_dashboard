#!/usr/bin/env python3
"""Extract versioned DPM template axes and resolve module taxonomies by date.

The EBA annotated-layout archives are release deltas in many versions. This
builder records each observed layout with source provenance, then derives
per-template intervals from the curated module_release_schedule.csv. Unknown
effective dates are retained as unresolved rows instead of being guessed.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import functools
import io
import json
import re
import sys
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "eba-dpm-history"
SOURCE_DIR = DATA / "sources"
OUTPUT_DIR = DATA / "generated"
SCHEDULE = DATA / "module_release_schedule.csv"

AXIS_COORDINATES = {
    "x": "x_axis_rc_code",
    "y": "y_axis_rc_code",
    "z": "z_axis_rc_code",
}
VERSION_RE = re.compile(r"(?<!\d)(4\.2\.1|4\.2|4\.1|4\.0|3\.5|3\.4|3\.3|3\.2|3\.1|3\.0(?:\.1)?|2\.10(?:\.1)?|2\.9\.1\.1)(?!\d)")
DOMAIN_REF_RE = re.compile(r"\(([A-Za-z0-9_]+):([A-Za-z0-9_]+)(?:\(([A-Za-z0-9_]+)\))?\)")
NUMERIC_CODE_RE = re.compile(r"^\d{1,4}$")
DISPLAY_PERCENT_RE = re.compile(r"(?i)(?:%|\bpercent(?:age)?\b|\bratio\b|\bpd\b|\blgd\b|\bocr\b|\bp2g\b)")
DISPLAY_RATE_RE = re.compile(r"(?i)\b(?:default|loss|cure|recovery|capital\s+buffer)\s+rate\b")
DISPLAY_COUNT_RE = re.compile(r"(?i)\b(?:number\s+of|count\s+of|number|count)\b")
DISPLAY_DURATION_RE = re.compile(r"(?i)\b(?:maturity|duration|repricing\s+time|survival\s+period)\b")
DISPLAY_DURATION_UNIT_RE = re.compile(r"(?i)\b(?:days?|months?|years?)\b")
REGULATORY_ACRONYMS = set("""
ABCP ACPR ACTP AVA BAFIN BIC BNRO BRRD BSI BTAR CCF CET1 CDS CNMV COREP CRD CRR
CSR CQS CSSF CVA CYSEC DGS DORA EAD EBA ECB EEPE ERBA ESG FINFSA FINREP FKTK FMI
FRTB FSMA FTNET FX GIRR GMRA GSIB HANFA HQLA ICMA IFRS IRBA IRB IRC ISDA ISIN
JST KDPW LCR LEI LGD LOCOM MFSA MREL NPE NPL NSFR OENB OCI OCR P2G PD PV QCCP
QRT RWA RWEA SEPA SFT SME SREP SWIFT TLAC TREA TPD TSCR UCITS XBRL
""".split())
COMMON_SHORT_WORDS = set("""
A AN AND ARE AS AT BE BEEN BUT BY CAN FOR FROM HAD HAS HAVE HE HER HIS HOW I IF IN
ALL ANY DAY DUE INTO IS IT ITS LOW MAY NET NEW NO NON NOT OF OFF OLD ON OR OUR
OWN OUT OVER PER PRE RAW ROW SHE SO TAB TOP THAN THAT THE THEIR THEM THEN THERE
THESE THEY THIS THOSE THROUGH TO UNDER UP VIA WAS WE WERE WHAT WHEN WHERE WHICH
WHO WHY WILL WITH WITHIN WITHOUT YOU YOUR AU AUX CE CES DANS DE DES DU ELLE EN ET
LA LE LES MAIS NE NI OU PAR PAS POUR QUE QUI SA SANS SE SES SON SUR UN UNE VERS
""".split())
ALL_CAPS_WORD_RE = re.compile(r"[\w]+", re.UNICODE)

# A curated display subset of the EBA's open currency domain. Keep the
# reporting currency first, then the major international currencies and the
# European/regional currencies most useful when browsing bank disclosures.
# The source glossary remains untouched; this order is applied on every build.
CURRENCY_Z_DISPLAY_ORDER = (
    "EUR", "USD", "GBP", "CHF", "JPY", "CNY",
    "CAD", "AUD", "SEK", "DKK", "NOK", "PLN",
    "CZK", "HUF", "RON", "SGD", "HKD", "TRY",
)

C08_IRB_Z_LABELS = {
    "with": "IRB A — With own estimates of LGD or conversion factors",
    "without": "IRB F — Without own estimates of LGD or conversion factors",
}
C08_IRB_QUALIFIER_RE = re.compile(
    r"\s*(?:[-–—]\s*)?(?:with|without) own estimates? of LGD (?:and/)?or conversion factors\s*$",
    re.IGNORECASE,
)
C08_IRB_PORTFOLIO_FAMILIES = (
    ("central governments", "central banks"),
    ("regional governments", "local authorities"),
    ("public sector entities",),
    ("institutions",),
    ("corporates",),
    ("retail",),
)


def limit_currency_z_rows(rows: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """Shorten open currency axes while retaining non-currency Z members.

    Native aggregate/other options have q-codes and remain available. Closed
    currency domains with at most 20 currencies keep all their members.
    """
    currency_rows = {code: label for code, label in rows if re.fullmatch(r"[A-Z]{3}", code)}
    if "EUR" not in currency_rows:
        return rows

    selected = [(code, currency_rows[code]) for code in CURRENCY_Z_DISPLAY_ORDER if code in currency_rows]
    if len(currency_rows) <= 20:
        selected.extend((code, label) for code, label in rows
                        if code in currency_rows and code not in CURRENCY_Z_DISPLAY_ORDER)
    special = [(code, label) for code, label in rows if not re.fullmatch(r"[A-Z]{3}", code)]
    aggregate = [(code, label) for code, label in special if "all currencies" in label.casefold()]
    other = [(code, label) for code, label in special if "all currencies" not in label.casefold()]
    return aggregate + selected + other


def c08_irb_portfolio_sort_key(description: str) -> tuple[bool, int]:
    """Prioritize major exposure families; keep memo items after them all."""
    label = description.casefold()
    memo = re.match(r"memo(?:randum)? items?\b\s*:?\s*", label)
    portfolio = label[memo.end():] if memo else label
    for rank, prefixes in enumerate(C08_IRB_PORTFOLIO_FAMILIES):
        if portfolio.startswith(prefixes):
            return bool(memo), rank
    return bool(memo), len(C08_IRB_PORTFOLIO_FAMILIES)


def group_c08_irb_z_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Use each real C_08 IRB total as the parent of its Z portfolios.

    Legacy releases place one Z choice on each numbered worksheet. Grouping
    the finished table/framework is necessary to assemble both complete IRB
    branches. The source total codes remain selectable; redundant qualifiers
    are removed from child labels without changing their codes or formats.
    """
    buckets: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        if row["coordinate"] == AXIS_COORDINATES["z"] and re.fullmatch(r"C_08\.\d{2}(?:\.\d+)?", row["table_id"]):
            buckets[(row["module_code"], row["framework"], row["table_id"])].append(row)

    replacements: dict[tuple[str, str, str], list[dict[str, str]]] = {}
    for key, members in buckets.items():
        groups: dict[str, list[dict[str, str]]] = {"with": [], "without": [], "other": []}
        for member in members:
            label = member["description"].casefold()
            if re.search(r"\bwithout own estimates? of lgd\b", label):
                groups["without"].append(member)
            elif re.search(r"\bwith own estimates? of lgd\b", label):
                groups["with"].append(member)
            else:
                groups["other"].append(member)
        if not groups["with"] or not groups["without"]:
            continue

        organized = [member.copy() for member in groups["other"]
                     if "all exposure classes and approaches" in member["description"].casefold()]
        for group_name in ("with", "without"):
            group_members = groups[group_name]
            totals = [member for member in group_members if member["description"].casefold().startswith("total ")]
            if len(totals) != 1:
                raise ValueError(f"Expected one C_08 IRB {group_name} total in {key}, found {len(totals)}")
            total = totals[0]
            organized.append({**total, "description": C08_IRB_Z_LABELS[group_name], "parent_coordinate_code": ""})
            portfolios = []
            for member in group_members:
                if member is total:
                    continue
                description = C08_IRB_QUALIFIER_RE.sub("", member["description"]).rstrip(" -–—")
                if not description or description == member["description"]:
                    raise ValueError(f"Cannot shorten C_08 IRB portfolio label: {key} {member['code']}")
                portfolios.append({**member, "description": description, "parent_coordinate_code": total["code"]})
            organized.extend(sorted(portfolios, key=lambda member: c08_irb_portfolio_sort_key(member["description"])))
        organized.extend(member.copy() for member in groups["other"]
                         if "all exposure classes and approaches" not in member["description"].casefold())
        for order, member in enumerate(organized, 1):
            member["order_first"] = str(order)
        replacements[key] = organized

    output: list[dict[str, str]] = []
    emitted: set[tuple[str, str, str]] = set()
    for row in rows:
        key = (row["module_code"], row["framework"], row["table_id"])
        if row["coordinate"] == AXIS_COORDINATES["z"] and key in replacements:
            if key not in emitted:
                output.extend(replacements[key])
                emitted.add(key)
            continue
        output.append(row)
    return output


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def format_taxonomy_label(value: Any) -> str:
    """Sentence-case shouting-style labels while preserving regulatory acronyms.

    This is applied to the extracted dictionary once, so the application only
    needs to display the dictionary's description values.
    """
    source = str(value or "")
    letters = [character for character in source if character.isalpha()]
    if not letters or any(character != character.upper() for character in letters):
        return source

    def format_word(match: re.Match[str]) -> str:
        word = match.group(0)
        upper = word.upper()
        plural_stem = upper[:-1] if upper.endswith("S") else ""
        if upper in REGULATORY_ACRONYMS or any(character.isdigit() for character in upper):
            return upper
        if plural_stem in REGULATORY_ACRONYMS:
            return f"{plural_stem}s"
        if len(upper) <= 3 and upper not in COMMON_SHORT_WORDS:
            return upper
        return word.lower()

    result = ALL_CAPS_WORD_RE.sub(format_word, source.lower())
    # Short country/reporting codes such as "IS" can also be ordinary English
    # words. Preserve them when the whole description consists of that code.
    if re.fullmatch(r"\s*[\w]+\s*", source, re.UNICODE) and len(source.strip()) == 2:
        result = source.strip()
    for index, character in enumerate(result):
        if character.isalpha():
            result = result[:index] + character.upper() + result[index + 1:]
            break
    return result


def code_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return f"{value:04d}" if 0 <= value <= 999 else str(value)
    if isinstance(value, float) and value.is_integer():
        return code_text(int(value))
    result = clean(value)
    if not NUMERIC_CODE_RE.fullmatch(result):
        return ""
    return str(int(result)) if len(result) > 1 and result.startswith("0") else result


def infer_display_format(description: str) -> str:
    """Infer display-only exceptions to the default monetary scale.

    The default is intentionally blank: ordinary monetary values continue to
    follow the application's selected currency scale. This V1 only tags
    percentage metrics and values that should remain in raw units.
    """
    segments = [clean(part) for part in str(description or "").split("/") if clean(part)]
    if not segments:
        return ""

    folded = [part.casefold() for part in segments]
    full_text = " / ".join(segments)
    leaf = segments[-1].strip()

    # Ratio descendants such as surplus/deficit, and explicit ratio
    # numerators or denominators, are amounts rather than percentage values.
    amount_component = re.compile(r"(?i)\b(?:numerator|denominator|surplus|deficit)\b")
    if amount_component.search(leaf):
        return ""

    # A percentage printed as an axis category (e.g. conversion factor 10%)
    # describes the category, not the scale of the cell value.
    if any("conversion factor" in part for part in folded):
        return ""
    if re.fullmatch(r"[+\-]?\d+(?:[.,]\d+)?\s*%", leaf):
        return ""

    if DISPLAY_PERCENT_RE.search(full_text) or DISPLAY_RATE_RE.search(full_text):
        return "%"

    if DISPLAY_COUNT_RE.search(full_text):
        return "Unit"

    if DISPLAY_DURATION_RE.search(full_text) and DISPLAY_DURATION_UNIT_RE.search(full_text):
        return "Unit"

    return ""


def normalize_template(value: Any) -> str:
    text = clean(value)
    text = re.sub(r"\s*\([^)]*\)\s*$", "", text)
    text = text.strip().replace(" ", "_")
    return text


def split_template_suffix(template_id: str) -> tuple[str, str] | None:
    """Split a final one-letter sub-template suffix from a template code."""
    match = re.fullmatch(r"(.+)\.([A-Za-z])", template_id.strip())
    return (match.group(1), match.group(2).casefold()) if match else None


def description_signature(value: str) -> str:
    return " ".join(value.casefold().split())


@functools.lru_cache(maxsize=1)
def production_display_format_lookup() -> dict[tuple[str, str, str, str], str]:
    """Load already-curated percent/raw-unit formats from the live dictionary."""
    source = ROOT / "app" / "assets" / "ITS_all_dimension_mapping.csv"
    if not source.exists():
        return {}
    values: dict[tuple[str, str, str, str], set[str]] = defaultdict(set)
    with source.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            value_format = clean(row.get("format", ""))
            if value_format not in {"%", "Unit"}:
                continue
            key = (
                clean(row.get("table_id", "")),
                clean(row.get("coordinate", "")),
                clean(row.get("code", "")),
                description_signature(clean(row.get("description", ""))),
            )
            values[key].add(value_format)
    return {key: next(iter(formats)) for key, formats in values.items() if len(formats) == 1}


def production_display_format(table_id: str, coordinate: str, code: str, description: str) -> str:
    key = (
        clean(table_id),
        clean(coordinate),
        clean(code),
        description_signature(clean(description)),
    )
    return production_display_format_lookup().get(key, "")


def dimension_display_format(table_id: str, coordinate: str, code: str, description: str) -> str:
    """Resolve the scale from the reported measure before label heuristics.

    C_80.00 has the same maturity labels under both Amount and Applicable RSF
    factor. In C_72–C_75, weight columns are fractions while caps mentioned
    in amount labels are only categories. C_76's sole X column mixes amounts
    and a ratio, distinguished by Y code 30. These source-specific rules cover
    all sheet suffixes and frameworks before consulting old curated formats.
    C_03/C_04 monetary rows also override legacy percentage tags and percent
    words that describe a threshold or risk weight rather than the cell value.
    """
    template = re.sub(r"\.[A-Za-z]$", "", table_id).upper()
    if template == "C_80.00":
        if coordinate == AXIS_COORDINATES["y"]:
            return ""
        if coordinate == AXIS_COORDINATES["x"]:
            return "%" if description.casefold().startswith("applicable rsf factor/") else ""
    if template in {"C_72.00", "C_73.00", "C_74.00", "C_75.01", "C_76.00"}:
        if coordinate == AXIS_COORDINATES["y"]:
            return "%" if template == "C_76.00" and code_text(code) == "30" else ""
        if coordinate == AXIS_COORDINATES["x"]:
            measure = description.casefold().strip()
            if template == "C_74.00":
                return "%" if measure.startswith("applicable weight/") else ""
            if template == "C_76.00":
                return ""
            return "%" if measure in {"applicable weight", "weight/applicable weight", "standard weight"} else ""
    if coordinate == AXIS_COORDINATES["y"]:
        if template == "C_03.00" and code_text(code) == "220":
            return ""  # Surplus/deficit of CET1 capital is a monetary amount.
        if template == "C_04.00":
            # Risk weights and CET1 thresholds describe the assets/conditions,
            # not the scale of the reported amount. Older code 900 has a
            # different label; only the actual output-floor rate is a ratio.
            is_output_floor_rate = code_text(code) == "900" and description.casefold().endswith("output floor applied (%)")
            return "%" if is_output_floor_rate else ""
    return production_display_format(table_id, coordinate, code, description) or infer_display_format(description)


def suffix_collision_signature(module: str, base: str, value: str) -> str:
    """Normalize known equivalent labels before treating suffixes as distinct.

    In legacy COREP C_08.01, sheets .a and .b share the specialized-lending
    Y-axis rows. The only textual difference is the parent label
    "Specialized lending slotting criteria (b)" vs "...: total"; the actual
    row codes and meanings are the same. Sheet .b adds X-axis rows such as
    code 100 that do not exist in .a, so the suffix does not disambiguate data.
    """
    signature = description_signature(value)
    if module == "COREP" and base == "C_08.01":
        signature = signature.replace(
            "specialized lending slotting criteria (b)",
            "specialized lending slotting criteria: total",
        )
    return signature


def find_discriminating_suffix_templates(
    mapping_rows: Iterable[dict[str, str]],
    template_rows: Iterable[dict[str, str]],
) -> tuple[set[tuple[str, str]], dict[tuple[str, str], int]]:
    """Find template families where suffix removal would create ambiguity.

    A suffix is retained only when two siblings can contain the same complete
    dimensional key (same code on every axis) and at least one shared code has
    a different meaning. A difference on one axis alone is not a collision if
    another axis has disjoint codes. Case-only description differences are
    treated as equivalent.
    """
    dimensions: dict[tuple[str, str, str, str], dict[str, dict[str, set[tuple[str, str, str]]]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(set))
    )
    for row in mapping_rows:
        split = split_template_suffix(row["table_id"])
        if not split:
            continue
        base, suffix = split
        key = (row["module_code"], row["framework"], base, row["coordinate"])
        dimensions[key][suffix][row["code"]].add((
            suffix_collision_signature(row["module_code"], base, row["description"]),
            row.get("format", ""),
            row.get("ignore", ""),
        ))

    conflicts: dict[tuple[str, str], int] = defaultdict(int)
    by_family: dict[tuple[str, str, str], dict[str, dict[str, dict[str, set[tuple[str, str, str]]]]]] = defaultdict(
        lambda: defaultdict(dict)
    )
    for (module, framework, base, coordinate), by_suffix in dimensions.items():
        for suffix, codes in by_suffix.items():
            by_family[(module, framework, base)][suffix][coordinate] = codes

    for (module, _framework, base), by_suffix in by_family.items():
        suffixes = sorted(by_suffix)
        for index, left_suffix in enumerate(suffixes):
            left_axes = by_suffix[left_suffix]
            for right_suffix in suffixes[index + 1:]:
                right_axes = by_suffix[right_suffix]
                # Different dimensional shapes cannot share a complete key.
                if left_axes.keys() != right_axes.keys():
                    continue
                shared_codes: dict[str, set[str]] = {}
                for coordinate in left_axes:
                    shared_codes[coordinate] = left_axes[coordinate].keys() & right_axes[coordinate].keys()
                # If any axis has no common code, the two subtemplates' cells
                # cannot collide after the suffix is removed.
                if any(not codes for codes in shared_codes.values()):
                    continue
                different_meanings = 0
                for coordinate, codes in shared_codes.items():
                    for code in codes:
                        if left_axes[coordinate][code] != right_axes[coordinate][code]:
                            different_meanings += 1
                if different_meanings:
                    conflicts[(module, base)] += different_meanings

    return set(conflicts), dict(conflicts)


def release_from_name(filename: str, default: str) -> str:
    match = VERSION_RE.search(filename)
    return match.group(1) if match else default


def workbook_module(name: str, archive_path: str) -> str:
    """Get the module identifier from the EBA package path/name."""
    name_only = Path(name).stem
    # Legacy DPM 1.0 names encode module as the final hyphen-separated token.
    legacy = re.search(r"\d{3}(?:-P\d)?-([A-Z][A-Z0-9_-]+)(?:\s|$)", name_only)
    if legacy:
        return legacy.group(1).split("-")[-1]
    newer_legacy = re.search(r"\b3\.5-([A-Z][A-Z0-9_]+)(?:\s|$)", name_only)
    if newer_legacy:
        return newer_legacy.group(1)

    # Newer package paths carry a reporting-area folder and a module code
    # concatenated with the area code (e.g. FINREP9FINREP, SEPA_IPRPAY).
    parts = Path(archive_path).parts
    area = next((p for p in reversed(parts[:-1]) if p.upper() in {
        "AE", "COREP", "DORA", "ESG", "FC", "FINREP", "FP", "GSII", "IF",
        "IMPRAC", "IPU", "IRRBB", "MICA", "MREL", "PAY", "PILLAR3", "REM", "RES", "SBP",
    }), "")
    matches = list(re.finditer(r"\b([A-Z][A-Z0-9_]*?" + re.escape(area) + r")\s+4\.\d", name_only)) if area else []
    if matches:
        token = matches[-1].group(1)
        return token[:-len(area)] or area
    # Flat 4.0/4.1 packages have no area directory. Strip date and version
    # markers, then split a module token from its reporting-area suffix.
    stem = re.sub(r"^\d{6,8}\s+", "", name_only)
    stem = re.sub(r"\s+4\.\d(?:\.\d)?(?:\s|$).*", "", stem)
    stem = re.sub(r"^Annotated Table Layout\s+", "", stem, flags=re.I).strip()
    for area_code in ("PILLAR3", "COREP", "FINREP", "MICA", "MREL", "DORA", "IRRBB", "IMPRAC", "GSII", "ESG", "PAY", "REM", "RES", "SBP", "AE", "FC", "FP", "IF", "IPU"):
        if stem.endswith(area_code):
            return stem[:-len(area_code)] or area_code
    return stem.split()[0] if stem else "UNKNOWN"


def sheet_template(sheet: Any) -> str:
    title = clean(sheet.cell(1, 1).value)
    if " - " in title:
        title = title.split(" - ", 1)[0]
    if not title:
        title = sheet.title
    return normalize_template(title)


def find_marker_rows(sheet: Any) -> tuple[int | None, int | None]:
    columns_row = rows_row = None
    for row in sheet.iter_rows():
        for cell in row[:6]:
            value = clean(cell.value).casefold()
            if value == "columns":
                columns_row = cell.row
            elif value == "rows":
                rows_row = cell.row
        if columns_row and rows_row:
            break
    return columns_row, rows_row


def with_parent_codes(
    entries: list[tuple[str, tuple[str, ...], int]],
) -> list[tuple[str, str, int, str, bool, str]]:
    """Add immediate code links and hidden nodes for uncoded header groups."""
    code_by_path: dict[tuple[str, ...], str] = {}
    for code, path, _row_no in entries:
        code_by_path.setdefault(path, code)

    virtual_by_path: dict[tuple[str, ...], str] = {}
    used_codes = {code for code, _path, _row_no in entries}

    def code_for_path(path: tuple[str, ...]) -> str:
        existing = code_by_path.get(path)
        if existing:
            return existing
        if path not in virtual_by_path:
            digest = hashlib.sha1("\u001f".join(path).encode("utf-8")).hexdigest()[:12]
            candidate = f"__PARENT__{digest}"
            salt = 1
            while candidate in used_codes:
                candidate = f"__PARENT__{digest}_{salt}"
                salt += 1
            used_codes.add(candidate)
            virtual_by_path[path] = candidate
        return virtual_by_path[path]

    for _code, path, _row_no in entries:
        for depth in range(1, len(path)):
            code_for_path(path[:depth])

    virtual_entries: list[tuple[str, str, int, str, bool, str]] = []
    for path, code in sorted(virtual_by_path.items(), key=lambda item: (len(item[0]), item[0])):
        parent_code = code_for_path(path[:-1]) if len(path) > 1 else ""
        virtual_entries.append((code, path[-1], 0, parent_code, True, "/".join(path)))

    real_entries = [
        (code, path[-1], row_no, code_for_path(path[:-1]) if len(path) > 1 else "", False, "/".join(path))
        for code, path, row_no in entries
    ]
    return virtual_entries + real_entries


def x_axis_rows(
    sheet: Any,
    columns_row: int | None,
    rows_row: int | None,
    *,
    include_parent: bool = False,
) -> list[tuple[str, str, int] | tuple[str, str, int, str, bool, str]]:
    if columns_row is None:
        return []
    end = rows_row or min(sheet.max_row, columns_row + 20)
    candidates: list[tuple[int, int]] = []
    for row_no in range(columns_row + 1, end):
        codes = sum(1 for col in range(1, sheet.max_column + 1) if code_text(sheet.cell(row_no, col).value))
        if codes:
            candidates.append((codes, row_no))
    if not candidates:
        return []
    max_count = max(count for count, _ in candidates)
    # Choose the last dense row of column codes; a few older layouts carry
    # auxiliary numeric identifiers above or below the actual code row.
    code_row = max(row for count, row in candidates if count == max_count)
    code_columns = [col for col in range(1, sheet.max_column + 1) if code_text(sheet.cell(code_row, col).value)]
    last_code_col = max(code_columns)

    # Treat each header row as a hierarchy level. A label remains active to
    # its right until another label appears at that level. When a new node
    # starts, deeper nodes from the previous branch are closed. This
    # deliberately ignores Excel merge ranges; anchors and blanks are handled
    # the same way as ordinary cells. Numeric members such as 0, 0.1 and 0.2
    # are retained as normal level labels.
    header_rows = range(columns_row + 1, code_row)
    active_by_level: dict[int, str] = {}
    output: list[tuple[str, tuple[str, ...], int]] = []
    for col in range(1, last_code_col + 1):
        for level, row_no in enumerate(header_rows):
            component = clean(sheet.cell(row_no, col).value)
            if not component:
                continue
            # Even identical text at a new cell begins a new node at this
            # level, so deeper labels cannot leak across sibling branches.
            active_by_level[level] = component
            for deeper_level in tuple(active_by_level):
                if deeper_level > level:
                    del active_by_level[deeper_level]

        code = code_text(sheet.cell(code_row, col).value)
        if not code:
            continue
        components = []
        for level in range(len(header_rows)):
            component = active_by_level.get(level, "")
            if component and (not components or components[-1] != component):
                components.append(component)
        if components:
            output.append((code, tuple(components), code_row))

    if not include_parent:
        return [(code, "/".join(components), row_no) for code, components, row_no in output]

    # Parentage comes from the reconstructed header levels, never from
    # splitting the display string: a component may itself contain "/".
    return with_parent_codes(output)


def y_axis_rows(
    sheet: Any,
    rows_row: int | None,
    table_id: str = "",
    *,
    include_parent: bool = False,
) -> list[tuple[str, str, int] | tuple[str, str, int, str, bool, str]]:
    if rows_row is None:
        return []
    items: list[tuple[str, str, int, int]] = []
    for row_no in range(rows_row, sheet.max_row + 1):
        code = code_text(sheet.cell(row_no, 3).value)
        label = clean(sheet.cell(row_no, 2).value)
        # A label without a selectable coordinate code can still be a real
        # hierarchy node. Keep it in the indentation stream so later coded
        # rows can inherit it as a parent; with_parent_codes() will assign an
        # internal synthetic code only when a descendant actually needs it.
        if not label or label.casefold() in {"code", "row", "rows"}:
            continue
        indent = sheet.cell(row_no, 2).alignment.indent or 0
        depth = max(0, round(float(indent) / 2))
        # In some FINREP F_18.00 releases, the source workbook accidentally
        # omits the single indentation level on code 5 (Cash balances). Keep
        # this correction scoped to that row/template, across all releases.
        if (
            code == "5"
            and label.casefold().startswith("cash balance")
            and re.fullmatch(r"F_18\.00(?:\.[a-z0-9]+)?", table_id, flags=re.IGNORECASE)
        ):
            depth = max(depth, 1)
        items.append((code or "", label, depth, row_no))

    if not items:
        return []

    # Most layouts list parents before their children. Use the absolute indent
    # as the level: equal indentation always closes the previous sibling.
    paths: list[list[str]] = []
    ancestors: list[tuple[int, str]] = []
    for _code, label, depth, _row_no in items:
        while ancestors and ancestors[-1][0] >= depth:
            ancestors.pop()
        paths.append([ancestor for _level, ancestor in ancestors] + [label])
        ancestors.append((depth, label))

    retroactive_moves: list[tuple[int, int]] = []

    def attach_preceding_block(start: int, end: int, parent_index: int) -> None:
        """Prefix a preceding deeper block with its later parent row."""
        parent_label = items[parent_index][1]
        block_ancestors: list[tuple[int, str]] = []
        for index in range(start, end):
            _code, label, depth, _row_no = items[index]
            while block_ancestors and block_ancestors[-1][0] >= depth:
                block_ancestors.pop()
            relative_path = [ancestor for _level, ancestor in block_ancestors] + [label]
            paths[index] = [parent_label] + relative_path
            block_ancestors.append((depth, label))
        paths[parent_index] = [parent_label]
        if start < end:
            retroactive_moves.append((items[parent_index][3], items[start][3]))

    # Retroactive parent detection applies to coded rows only. Some layouts
    # append uncoded notes after the final coded row, and those notes must not
    # displace a genuine trailing total (e.g. F_18.00's code 550). Likewise, a
    # leading uncoded section title is already the root of its descendants and
    # prevents a later top-level coded sibling from stealing that block.
    coded_indices = [index for index, item in enumerate(items) if item[0]]
    first_minimum = -1
    if coded_indices:
        minimum_depth = min(items[index][2] for index in coded_indices)
        first_minimum = next(index for index in coded_indices if items[index][2] == minimum_depth)
        preceding_coded = [index for index in coded_indices if index < first_minimum]
        preceding_uncoded_root = any(
            not items[index][0] and items[index][2] <= minimum_depth
            for index in range(first_minimum)
        )
        if (
            preceding_coded
            and not preceding_uncoded_root
            and all(items[index][2] > minimum_depth for index in preceding_coded)
        ):
            attach_preceding_block(0, first_minimum, first_minimum)

    # A final coded, less-indented total can parent the trailing deeper block.
    # Ignore later uncoded notes when choosing that terminal row, and stop at
    # the previous coded row at the same or a shallower indentation.
    if len(coded_indices) > 1:
        parent_index = coded_indices[-1]
        previous_coded_index = coded_indices[-2]
        parent_depth = items[parent_index][2]
        start = previous_coded_index
        while start >= 0:
            if items[start][0] and items[start][2] <= parent_depth:
                break
            start -= 1
        start += 1
        preceding_uncoded_root = any(
            not items[index][0] and items[index][2] <= parent_depth
            for index in range(start, parent_index)
        )
        if (
            items[previous_coded_index][2] > parent_depth
            and parent_index != first_minimum
            and not preceding_uncoded_root
        ):
            if start < parent_index:
                attach_preceding_block(start, parent_index, parent_index)

    # Keep the source order except where a parent was identified after its
    # children: move it immediately before the first child it now parents.
    ordered = list(zip(items, paths))
    for parent_row_no, first_child_row_no in retroactive_moves:
        parent_index = next(index for index, (item, _path) in enumerate(ordered) if item[3] == parent_row_no)
        parent = ordered.pop(parent_index)
        child_index = next(index for index, (item, _path) in enumerate(ordered) if item[3] == first_child_row_no)
        ordered.insert(child_index, parent)

    if not include_parent:
        return [(item[0], "/".join(path), item[3]) for item, path in ordered if item[0]]

    return with_parent_codes([(item[0], tuple(path), item[3]) for item, path in ordered if item[0]])


@functools.lru_cache(maxsize=32)
def dictionary_items(dictionary_name: str) -> dict[str, list[tuple[str, str]]]:
    dictionary_path = Path(dictionary_name)
    if not dictionary_path.exists():
        return {}
    workbook = load_workbook(dictionary_path, read_only=True, data_only=True)
    result: dict[str, list[tuple[str, str]]] = defaultdict(list)
    try:
        if "Members" in workbook.sheetnames:
            sheet = workbook["Members"]
            header = [clean(c.value) for c in next(sheet.iter_rows(min_row=1, max_row=1))]
            idx = {name: i for i, name in enumerate(header) if name}
            if {"DomainCode", "MemberCode", "MemberLabel"}.issubset(idx):
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    domain = clean(row[idx["DomainCode"]])
                    code = clean(row[idx["MemberCode"]])
                    label = clean(row[idx["MemberLabel"]])
                    if domain and code and label:
                        result[domain].append(("q" + code if code.startswith("x") else code, label))
        elif "Item" in workbook.sheetnames:
            sheet = workbook["Item"]
            header = [clean(c.value) for c in next(sheet.iter_rows(min_row=1, max_row=1))]
            idx = {name: i for i, name in enumerate(header) if name}
            if {"CategoryCode", "ItemCode", "Name"}.issubset(idx):
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    category = clean(row[idx["CategoryCode"]])
                    code = clean(row[idx["ItemCode"]])
                    label = clean(row[idx["Name"]])
                    if category and code and label:
                        result[category].append((code, label))

        # Sheet-per dimensions can point to a subcategory rather than the
        # entire category. Keep these member sets separately so annotations
        # such as ``(qEC:qEC2) <Key value>`` do not expand to every qEC value.
        for sheet_name in ("SubCategoryItems", "SubCategoryItem"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            header = [clean(c.value) for c in next(sheet.iter_rows(min_row=1, max_row=1))]
            idx = {name: i for i, name in enumerate(header) if name}
            category_column = "CategoryCode" if "CategoryCode" in idx else "Category_SubCategory"
            required = {category_column, "SubCategoryCode", "ItemCode", "ItemName"}
            if not required.issubset(idx):
                continue
            for row in sheet.iter_rows(min_row=2, values_only=True):
                category = clean(row[idx[category_column]])
                subcategory = clean(row[idx["SubCategoryCode"]])
                code = clean(row[idx["ItemCode"]])
                label = clean(row[idx["ItemName"]])
                if category and subcategory and code and label:
                    code = "q" + code if code.startswith("x") else code
                    members = result[f"{category}:{subcategory}"]
                    if not any(existing_code == code for existing_code, _ in members):
                        members.append((code, label))
    finally:
        workbook.close()
    return dict(result)


def dictionary_z_rows(dictionary_path: Path, sheet: Any) -> list[tuple[str, str]]:
    """Resolve enumerated sheet-per/key-value dimensions from DPM glossaries."""
    if not dictionary_path.exists():
        return []
    refs: list[tuple[str, str]] = []
    key_value_refs: list[tuple[str, str]] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 5)):
        for cell in row:
            text = clean(cell.value)
            if not text:
                continue
            for match in DOMAIN_REF_RE.finditer(text):
                first, second, nested = match.groups()
                # DPM 2.0 marks a tab dimension with <Key value>. Older DPM
                # annotates a varying dimension as (dimension:domain), while
                # fixed selections use member codes such as (AP:x42).
                if "key value" in text.casefold():
                    # In the older annotation ``(qEC:qEC2) <Key value>``,
                    # the first token is the dimension domain and the second
                    # is its subdomain. Newer annotations add the domain again
                    # as a nested token: ``(qEC:qEC2(qEC))``. Prefer that
                    # explicit nested domain, otherwise use the first token.
                    key_value_refs.append((nested or first, second))
                elif re.fullmatch(r"[A-Z]{2,5}", second):
                    refs.append((second, second))
    is_key_value = bool(key_value_refs)
    refs = list(dict.fromkeys(key_value_refs or refs))
    if not refs:
        return []

    all_items = dictionary_items(str(dictionary_path.resolve()))
    items: list[tuple[str, str]] = []
    for category_or_domain, subcode in refs:
        if is_key_value:
            # Use the exact referenced subcategory when available. Some DPM
            # releases only provide a category reference, so keep the full
            # category as a compatibility fallback for those annotations.
            items.extend(
                all_items.get(f"{category_or_domain}:{subcode}")
                or all_items.get(category_or_domain, [])
            )
        else:
            items.extend(all_items.get(category_or_domain, []))
    unique: dict[str, str] = {}
    for code, label in items:
        unique.setdefault(code, label)
    return list(unique.items())


def legacy_sheet_z_rows(sheet: Any) -> list[tuple[str, str]]:
    """Read a DPM 1.0 sheet-per Z member from its numbered layout sheet.

    Legacy annotated packages materialise each tab value as its own worksheet:
    the sheet title ends in ``(001)``, and the matching label is repeated in
    the first rows. Expanding every referenced glossary domain here would
    incorrectly add hundreds of unrelated members to the template's Z axis.
    """
    match = re.search(r"\((\d{3,4})\)\s*$", clean(sheet.title))
    if not match:
        return []
    code = match.group(1)
    numeric_code = int(code)

    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 5)):
        for cell in row:
            text = clean(cell.value)
            label_match = re.match(r"0*(\d{1,4})\s+(.+)$", text)
            if label_match and int(label_match.group(1)) == numeric_code:
                return [(code, label_match.group(2))]

    return []


def read_schedule() -> list[dict[str, str]]:
    with SCHEDULE.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def schedule_match(module: str, framework: str, template_id: str, rows: list[dict[str, str]]) -> dict[str, str] | None:
    normalized = module.upper()
    family = ".".join(framework.split(".")[:2])
    eligible_versions = {framework, family}
    if family == "3.0":
        eligible_versions.add("3.0.1")
    candidates: list[tuple[tuple[int, int], dict[str, str]]] = []
    for row in rows:
        pattern = row["module_pattern"].upper()
        if pattern == "*" or (pattern.endswith("*") and normalized.startswith(pattern[:-1])) or pattern == normalized:
            template_pattern = (row.get("template_pattern") or "").upper()
            if template_pattern and not template_id.upper().startswith(template_pattern):
                continue
            # Old packages use broad taxonomy IDs; only rows for the exact
            # framework, plus explicit earlier base rows, are applicable.
            if row["framework"] in eligible_versions:
                specificity = (len(pattern.rstrip("*")), len(template_pattern))
                candidates.append((specificity, row))
    return max(candidates, key=lambda item: item[0])[1] if candidates else None


def close_template_intervals(template_rows: Iterable[dict[str, str]]) -> list[dict[str, str]]:
    """Close snapshots only within the same module/template identity.

    A template code can be reused by unrelated modules, and some releases
    split one module into several parallel DPM 2.0 modules. Cross-module
    interval closure would incorrectly make one branch truncate another.
    Module-family aliases are resolved separately by resolve_taxonomy.py.
    """
    ordered = sorted(template_rows, key=lambda r: (r["module_code"], r["template_id"], r["effective_from"] or "9999", r["framework"]))
    by_module_template: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in ordered:
        if row["effective_from"]:
            by_module_template[(row["module_code"], row["template_id"])].append(row)
    for versions in by_module_template.values():
        versions.sort(key=lambda r: (r["effective_from"], r["framework"]))
        for current, following in zip(versions, versions[1:]):
            current["effective_to"] = following["effective_from"]
    return ordered


def extract_all() -> tuple[
    list[dict[str, str]],
    list[dict[str, str]],
    list[dict[str, str]],
    dict[tuple[str, str], int],
]:
    schedules = read_schedule()
    layout_rows: list[dict[str, str]] = []
    template_rows: dict[tuple[str, str, str], dict[str, str]] = {}
    inventory: list[dict[str, str]] = []
    dictionaries = {p.stem.removesuffix("_dictionary"): p for p in SOURCE_DIR.glob("*_dictionary.xlsx")}

    for archive in sorted(SOURCE_DIR.glob("*_layouts.zip")):
        archive_framework = archive.name.removesuffix("_layouts.zip")
        with zipfile.ZipFile(archive) as bundle:
            workbook_names = [n for n in bundle.namelist() if n.lower().endswith(".xlsx")]
            for member in workbook_names:
                module = workbook_module(member, member)
                version = release_from_name(Path(member).name, archive_framework)
                dictionary = dictionaries.get(version) or dictionaries.get(archive_framework)
                try:
                    workbook = load_workbook(io.BytesIO(bundle.read(member)), data_only=True)
                except Exception as error:
                    inventory.append({"framework": version, "module": module, "workbook": member, "sheets": "", "status": f"workbook_error: {error}"})
                    continue
                table_count = 0
                candidate_sheets = [sheet for sheet in workbook.worksheets if sheet.title.casefold() not in {"toc", "contents", "readme", "index"}]
                table_sheet_counts = defaultdict(int)
                for candidate in candidate_sheets:
                    table_sheet_counts[sheet_template(candidate)] += 1
                for sheet in candidate_sheets:
                    table_id = sheet_template(sheet)
                    if not table_id or table_id.casefold() in {"table of contents", "toc"}:
                        continue
                    columns_row, rows_row = find_marker_rows(sheet)
                    xs = x_axis_rows(sheet, columns_row, rows_row, include_parent=True)
                    ys = y_axis_rows(sheet, rows_row, table_id, include_parent=True)
                    header_text = " ".join(
                        clean(sheet.cell(r, c).value)
                        for r in range(1, min(sheet.max_row, 5) + 1)
                        for c in range(1, min(sheet.max_column, 12) + 1)
                    ).casefold()
                    if version.startswith("4."):
                        # DPM 2.0 uses Key value for many cell-level component
                        # dimensions. Only the explicit sheet-per declaration
                        # denotes the table's Z axis.
                        has_variable_tab_axis = "sheet per" in header_text
                    else:
                        has_variable_tab_axis = (
                            "sheet per" in header_text
                            or (table_sheet_counts[table_id] > 1 and bool(re.search(r"\([^)]*\d{3,4}\)\s*$", clean(sheet.title))))
                        )
                    if has_variable_tab_axis:
                        # DPM 1.0 uses one numbered worksheet per tab/Z member;
                        # DPM 2.0 declares enumerated tab dimensions in the
                        # glossary via a Key value reference.
                        if version.startswith("4."):
                            zs = [(code, desc, 0) for code, desc in limit_currency_z_rows(dictionary_z_rows(dictionary, sheet))]
                        else:
                            zs = [(code, desc, int(code)) for code, desc in legacy_sheet_z_rows(sheet)]
                    else:
                        zs = []
                    effective = schedule_match(module, version, table_id, schedules)
                    status = effective["status"] if effective else "needs_effective_date"
                    from_date = effective["effective_from"] if effective else ""
                    source_note = effective["evidence"] if effective else "No module/date rule matched; do not infer an effective date."
                    if columns_row is None and rows_row is None:
                        continue
                    table_count += 1
                    template_key = (module, table_id, version)
                    template_rows[template_key] = {
                        "module_code": module,
                        "template_id": table_id,
                        "framework": version,
                        "effective_from": from_date,
                        "effective_to": "",
                        "status": status,
                        "effective_source": effective["source_url"] if effective else "",
                        "effective_evidence": source_note,
                        "source_archive": archive.name,
                        "source_workbook": member,
                        "source_sheet": sheet.title,
                    }
                    for axis, entries in (("x", xs), ("y", ys), ("z", zs)):
                        for order, entry in enumerate(entries, 1):
                            code, description, _row_number = entry[:3]
                            parent_code = entry[3] if len(entry) > 3 else ""
                            is_virtual_parent = bool(entry[4]) if len(entry) > 4 else False
                            format_description = entry[5] if len(entry) > 5 else description
                            layout_rows.append({
                                "table_id": table_id,
                                "coordinate": AXIS_COORDINATES[axis],
                                "code": code,
                                "parent_coordinate_code": parent_code,
                                "description": format_taxonomy_label(description),
                                "order_first": "" if axis == "y" else str(_row_number if axis == "z" and _row_number else order),
                                "ignore": "Y" if is_virtual_parent else "",
                                "format": "" if is_virtual_parent else dimension_display_format(
                                    table_id, AXIS_COORDINATES[axis], code, format_description
                                ),
                                "module_code": module,
                                "framework": version,
                                "effective_from": from_date,
                                "effective_to": "",
                                "status": status,
                                "source_workbook": member,
                                "source_sheet": sheet.title,
                                "effective_source": effective["source_url"] if effective else "",
                            })
                workbook.close()
                inventory.append({
                    "framework": version,
                    "module": module,
                    "workbook": member,
                    "sheets": str(table_count),
                    "status": "parsed",
                })

    discriminating, suffix_conflicts = find_discriminating_suffix_templates(
        [row for row in layout_rows if row.get("ignore") != "Y"],
        template_rows.values(),
    )

    # Suffixes are usually display/segmentation variants of one template. Keep
    # them only for module/template families where stripping them would make a
    # reused axis code ambiguous.
    for row in layout_rows:
        split = split_template_suffix(row["table_id"])
        if split and (row["module_code"], split[0]) not in discriminating:
            row["table_id"] = split[0]

    normalized_templates: dict[tuple[str, str, str], dict[str, str]] = {}
    for row in template_rows.values():
        split = split_template_suffix(row["template_id"])
        normalized_id = row["template_id"]
        if split and (row["module_code"], split[0]) not in discriminating:
            normalized_id = split[0]
        normalized = {**row, "template_id": normalized_id}
        normalized_templates.setdefault((row["module_code"], normalized_id, row["framework"]), normalized)

    # Collapse duplicate sheet tabs, retaining the first occurrence for an
    # identical table/axis/code/release and preserving full source inventory.
    unique_layouts: dict[tuple[str, str, str, str, str], dict[str, str]] = {}
    for row in layout_rows:
        key = (row["module_code"], row["framework"], row["table_id"], row["coordinate"], row["code"])
        unique_layouts.setdefault(key, row)

    # Materialise intervals within each module/template. DPM 2.0 module splits
    # remain visible as separate module histories and are followed by the
    # resolver when the caller requests the former module family.
    sorted_templates = close_template_intervals(normalized_templates.values())

    dimensions = group_c08_irb_z_rows(list(unique_layouts.values()))
    return dimensions, sorted_templates, inventory, suffix_conflicts


def write_csv(path: Path, rows: Iterable[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    dimensions, templates, inventory, suffix_conflicts = extract_all()
    write_csv(args.output_dir / "versioned_dimension_mapping.csv", dimensions, [
        "table_id", "coordinate", "code", "parent_coordinate_code", "description", "order_first", "ignore", "format",
        "module_code", "framework", "effective_from", "effective_to", "status", "source_workbook", "source_sheet", "effective_source",
    ])
    write_csv(args.output_dir / "template_taxonomy_history.csv", templates, [
        "module_code", "template_id", "framework", "effective_from", "effective_to", "status",
        "effective_source", "effective_evidence", "source_archive", "source_workbook", "source_sheet",
    ])
    write_csv(args.output_dir / "workbook_inventory.csv", inventory, ["framework", "module", "workbook", "sheets", "status"])
    unresolved = sorted({
        r["module_code"] for r in templates
        if not r["effective_from"] and not r["status"].startswith(("not_applicable", "superseded"))
    })
    report = {
        "layout_rows": len(dimensions),
        "inferred_display_formats": {
            value: sum(1 for row in dimensions if row.get("format") == value)
            for value in ("%", "Unit")
        },
        "template_release_rows": len(templates),
        "workbooks": len(inventory),
        "workbooks_with_errors": sum(1 for r in inventory if r["status"].startswith("workbook_error")),
        "modules_without_effective_date_rule": unresolved,
        "schedule_rows": len(read_schedule()),
        "templates_with_discriminating_suffixes": len(suffix_conflicts),
        "unique_discriminating_template_ids": len({template for _module, template in suffix_conflicts}),
        "discriminating_suffix_templates": [
            {"module_code": module, "template_id": template, "conflicting_axis_code_cases": count}
            for (module, template), count in sorted(suffix_conflicts.items())
        ],
        "generated_from": "official EBA source packages listed in sources.json",
    }
    (args.output_dir / "build_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["workbooks_with_errors"] == 0 else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"taxonomy build failed: {error}", file=sys.stderr)
        raise
