#!/usr/bin/env python3
"""Build the app's canonical template-name catalog from EBA layout sources.

The versioned dimension mapping retains the source workbook and worksheet for
each extracted template. This script reads the worksheet's official title,
uses the newest extracted framework available for each app template ID, and
collapses presentation-only letter suffixes with the same rule as the
taxonomy extraction pipeline.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from build_taxonomy import format_taxonomy_label, release_from_name

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "eba-dpm-history"
SOURCE_DIR = DATA / "sources"
MAPPING_SOURCE = DATA / "generated" / "versioned_dimension_mapping.csv"
HISTORY_SOURCE = DATA / "generated" / "template_taxonomy_history.csv"
APP_MAPPING = ROOT / "app" / "assets" / "ITS_all_dimension_mapping.csv"
CSV_OUTPUT = ROOT / "app" / "assets" / "ITS_explorer_template_names.csv"
JS_OUTPUT = ROOT / "app" / "src" / "data" / "explorerTemplateNames.js"

CSV_COLUMNS = [
    "table_id",
    "description",
    "framework",
    "source_archive",
    "module_codes",
    "source_template_id",
    "source_workbook",
    "source_sheet",
]
SPLIT_TITLE_SUFFIX = re.compile(r"\s*\((?:[IVXLCDM]+|[a-z])\)\s*$", re.IGNORECASE)


def read_csv(path: Path, delimiter: str) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def framework_key(value: str) -> tuple[int, ...]:
    try:
        return tuple(int(part) for part in value.split("."))
    except ValueError as error:
        raise ValueError(f"Unexpected framework value: {value!r}") from error


def normalize_template(value: Any) -> str:
    text = " ".join(str(value or "").split())
    text = re.sub(r"\s*\([^)]*\)\s*$", "", text)
    return text.strip().replace(" ", "_")


def find_source_archive(framework: str, workbook_name: str, archive_members: dict[str, list[tuple[Path, str]]]) -> Path:
    candidates = archive_members.get(workbook_name, [])
    matching = [archive for archive, release in candidates if release == framework]
    if len(matching) == 1:
        return matching[0]
    if len(matching) > 1:
        raise ValueError(f"Multiple EBA archives contain {workbook_name!r} for framework {framework}.")
    raise ValueError(f"No EBA archive contains {workbook_name!r} for framework {framework}.")


def extract_sheet_title(sheet: Any, table_id: str) -> tuple[str, str]:
    """Return normalized source worksheet ID and its official title."""
    a1 = " ".join(str(sheet.cell(1, 1).value or "").split())
    if " - " in a1:
        raw_id, description = a1.split(" - ", 1)
        source_id = normalize_template(raw_id)
        return source_id, description.strip()

    source_id = normalize_template(a1) or normalize_template(sheet.title)
    toc = next((
        candidate for candidate in sheet.parent.worksheets
        if candidate.title.strip().casefold() in {"toc", "contents", "table of contents"}
    ), None)
    if toc:
        for row in toc.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            for index, value in enumerate(values):
                if normalize_template(value) == source_id:
                    description = next((item for item in values[index + 1:] if item), "")
                    if description:
                        return source_id, description

    # Some legacy sheets put the description in the first populated cell of
    # row one without an explicit "code - title" separator.
    for cell in sheet[1]:
        value = " ".join(str(cell.value or "").split())
        if value and normalize_template(value) != source_id and value != table_id:
            return source_id, value
    return source_id, ""


def source_variant_rank(record: dict[str, str], table_id: str) -> tuple[int, int, str, str, str]:
    source_id = record["source_template_id"]
    if source_id == table_id:
        return (0, 0, "", record["module_code"], record["source_sheet"])

    suffix = source_id[len(table_id) + 1:] if source_id.startswith(table_id + ".") else ""
    if len(suffix) == 1 and suffix.isalpha():
        return (1, 0 if suffix.casefold() == "a" else 1, suffix.casefold(), record["module_code"], record["source_sheet"])
    return (2, 0, source_id.casefold(), record["module_code"], record["source_sheet"])


def merged_variant_title(record: dict[str, str], table_id: str) -> str:
    """Remove a sheet-part marker only when its letter suffix was collapsed."""
    source_id = record["source_template_id"]
    suffix = source_id[len(table_id) + 1:] if source_id.startswith(table_id + ".") else ""
    if len(suffix) == 1 and suffix.isalpha():
        return SPLIT_TITLE_SUFFIX.sub("", record["description"]).strip()
    return record["description"].strip()


def get_source_records() -> list[dict[str, str]]:
    mapping_rows = read_csv(MAPPING_SOURCE, ",")
    history_rows = read_csv(HISTORY_SOURCE, ",")
    references: dict[tuple[str, str, str, str, str], dict[str, str]] = {}

    for row in mapping_rows:
        key = (
            row.get("table_id", "").strip(),
            row.get("module_code", "").strip(),
            row.get("framework", "").strip(),
            row.get("source_workbook", "").strip(),
            row.get("source_sheet", "").strip(),
        )
        if all(key):
            references.setdefault(key, {
                "table_id": key[0], "module_code": key[1], "framework": key[2],
                "source_workbook": key[3], "source_sheet": key[4],
            })

    # Some templates intentionally have no extractable axis rows. Their
    # taxonomy-history records still point to the official layout worksheet.
    for row in history_rows:
        key = (
            row.get("template_id", "").strip(),
            row.get("module_code", "").strip(),
            row.get("framework", "").strip(),
            row.get("source_workbook", "").strip(),
            row.get("source_sheet", "").strip(),
        )
        if all(key):
            references.setdefault(key, {
                "table_id": key[0], "module_code": key[1], "framework": key[2],
                "source_workbook": key[3], "source_sheet": key[4],
            })

    archive_members: dict[str, list[tuple[Path, str]]] = defaultdict(list)
    for archive in sorted(SOURCE_DIR.glob("*_layouts.zip")):
        default_framework = archive.name.removesuffix("_layouts.zip")
        with zipfile.ZipFile(archive) as bundle:
            for member in bundle.namelist():
                if member.lower().endswith(".xlsx"):
                    framework = release_from_name(Path(member).name, default_framework)
                    archive_members[member].append((archive, framework))

    workbooks: dict[tuple[Path, str], Any] = {}
    records: list[dict[str, str]] = []
    try:
        for reference in references.values():
            archive = find_source_archive(
                reference["framework"], reference["source_workbook"], archive_members
            )
            key = (archive, reference["source_workbook"])
            if key not in workbooks:
                with zipfile.ZipFile(archive) as bundle:
                    contents = bundle.read(reference["source_workbook"])
                workbooks[key] = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            workbook = workbooks[key]
            if reference["source_sheet"] not in workbook.sheetnames:
                raise ValueError(
                    f"Missing source worksheet {reference['source_sheet']!r} in "
                    f"{reference['source_workbook']!r}."
                )
            source_id, description = extract_sheet_title(
                workbook[reference["source_sheet"]], reference["table_id"]
            )
            if not description:
                raise ValueError(
                    f"No title in source worksheet {reference['source_sheet']!r} "
                    f"({reference['source_workbook']!r})."
                )
            records.append({
                **reference,
                "source_template_id": source_id,
                "description": " ".join(description.split()),
                "source_archive": archive.name,
            })
    finally:
        for workbook in workbooks.values():
            workbook.close()

    return records


def build_template_names() -> list[dict[str, str]]:
    app_ids = {
        row["table_id"].strip()
        for row in read_csv(APP_MAPPING, ";")
        if row.get("table_id", "").strip()
    }
    by_template: dict[str, list[dict[str, str]]] = defaultdict(list)
    for record in get_source_records():
        if record["table_id"] in app_ids:
            by_template[record["table_id"]].append(record)

    output: list[dict[str, str]] = []
    unresolved: list[str] = []
    for table_id in sorted(app_ids):
        records = by_template.get(table_id, [])
        if not records:
            unresolved.append(f"{table_id}: no extracted source record")
            continue

        latest_framework = max((record["framework"] for record in records), key=framework_key)
        current = [record for record in records if record["framework"] == latest_framework]
        by_module: dict[str, list[dict[str, str]]] = defaultdict(list)
        for record in current:
            by_module[record["module_code"]].append(record)

        chosen_by_module = []
        for module_code, variants in sorted(by_module.items()):
            chosen = min(variants, key=lambda record: source_variant_rank(record, table_id))
            chosen_by_module.append((module_code, chosen, merged_variant_title(chosen, table_id)))

        descriptions = {description for _module, _record, description in chosen_by_module}
        if len(descriptions) != 1:
            variants = "; ".join(
                f"{module}: {description}" for module, _record, description in chosen_by_module
            )
            unresolved.append(f"{table_id}: conflicting latest official titles ({variants})")
            continue

        description = format_taxonomy_label(next(iter(descriptions)))
        _source_module, source_record, _description = chosen_by_module[0]
        output.append({
            "table_id": table_id,
            "description": description,
            "framework": latest_framework,
            "source_archive": source_record["source_archive"],
            "module_codes": ",".join(module for module, _record, _label in chosen_by_module),
            "source_template_id": source_record["source_template_id"],
            "source_workbook": source_record["source_workbook"],
            "source_sheet": source_record["source_sheet"],
        })

    if unresolved:
        raise ValueError("Template-name catalog needs review:\n- " + "\n- ".join(unresolved))
    if len(output) != len(app_ids):
        raise ValueError(f"Expected {len(app_ids)} template names, generated {len(output)}.")

    output.sort(key=lambda row: row["table_id"].casefold())
    write_outputs(output)
    return output


def write_outputs(rows: list[dict[str, str]]) -> None:
    with CSV_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, delimiter=";", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    names = {row["table_id"]: row["description"] for row in rows}
    lines = [
        "// Generated from EBA annotated template layouts by scripts/dpm_taxonomy/build_template_names.py.",
        "export const EXPLORER_TEMPLATE_LABELS = Object.freeze({",
    ]
    lines.extend(f"  {json.dumps(table_id, ensure_ascii=False)}: {json.dumps(name, ensure_ascii=False)}," for table_id, name in names.items())
    lines.extend(["});", ""])
    JS_OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    try:
        rows = build_template_names()
    except Exception as error:  # noqa: BLE001 - command-line diagnostic should be concise
        print(f"ERROR: {error}", file=sys.stderr)
        return 1
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        counts[row["framework"]] += 1
    print(f"Wrote {len(rows):,} template names to {CSV_OUTPUT} and {JS_OUTPUT}")
    print("Latest-source framework counts: " + ", ".join(
        f"{framework}: {count}" for framework, count in sorted(counts.items(), key=lambda item: framework_key(item[0]))
    ))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
