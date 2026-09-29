import io
import sys
import unittest
import zipfile
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
import build_taxonomy as builder  # noqa: E402
from resolve_taxonomy import resolve  # noqa: E402


ROOT = SCRIPT_DIR.parents[1]
SOURCES = ROOT / "data" / "eba-dpm-history" / "sources"


class DpmTaxonomyTests(unittest.TestCase):
    def test_c08_irb_z_groups_legacy_and_qx_codes_without_changing_identities(self):
        def dimension(table_id, coordinate, code, description):
            return {
                "module_code": "COREP", "framework": "3.2", "table_id": table_id,
                "coordinate": coordinate, "code": code, "description": description,
                "parent_coordinate_code": "", "ignore": "", "format": "", "order_first": "",
            }

        original = [
            dimension("C_08.01", "z_axis_rc_code", "0002", "Total without own estimates of LGD or conversion factors"),
            dimension("C_08.01", "x_axis_rc_code", "0010", "Amount"),
            dimension("C_08.01", "z_axis_rc_code", "0003", "Central banks with own estimates of LGD or conversion factors"),
            dimension("C_08.01", "z_axis_rc_code", "0001", "Total with own estimates of LGD or conversion factors"),
            dimension("C_08.01", "z_axis_rc_code", "0004", "Central banks without own estimates of LGD or conversion factors"),
            dimension("C_08.01", "z_axis_rc_code", "qx01", "All exposure classes and approaches"),
            dimension("C_08.01", "z_axis_rc_code", "qx2082", "Collective Investment Undertakings (CIU)"),
            dimension("C_08.06", "z_axis_rc_code", "qx129", "Project finance"),
        ]
        result = builder.group_c08_irb_z_rows(original)
        grouped = [row for row in result if row["table_id"] == "C_08.01" and row["coordinate"] == "z_axis_rc_code"]
        self.assertEqual([row["code"] for row in grouped], [
            "qx01", "__PARENT__C08_IRB_A", "0001", "0003",
            "__PARENT__C08_IRB_F", "0002", "0004", "qx2082",
        ])
        self.assertEqual([row["parent_coordinate_code"] for row in grouped[2:4]], ["__PARENT__C08_IRB_A"] * 2)
        self.assertEqual([row["parent_coordinate_code"] for row in grouped[5:7]], ["__PARENT__C08_IRB_F"] * 2)
        self.assertEqual([row["ignore"] for row in (grouped[1], grouped[4])], ["Y", "Y"])
        self.assertEqual([row["order_first"] for row in grouped], [str(number) for number in range(1, 9)])
        self.assertEqual(next(row for row in result if row["table_id"] == "C_08.06"), original[-1])
        self.assertEqual(original[0]["parent_coordinate_code"], "")

    def test_c03_c04_amounts_override_percent_words_and_stale_formats(self):
        cases = {
            ("C_03.00", "220", "Surplus(+)/Deficit(-) of CET1 capital considering OCR and P2G"): "",
            ("C_03.00", "130", "Total SREP capital requirement ratio (TSCR)"): "%",
            ("C_04.00", "96", "Deferred tax assets subject to a risk weight of 250%"): "",
            ("C_04.00", "97", "Deferred tax assets subject to a risk weight of 0%"): "",
            ("C_04.00", "200", "10% CET1 threshold"): "",
            ("C_04.00", "210", "17.65% CET1 threshold"): "",
            ("C_04.00", "504", "Investments in CET1 capital - subject to a risk weight of 250%"): "",
            ("C_04.00", "900", "Output floor/Output floor applied (%)"): "%",
            ("C_04.00", "900", "Basel I floor/Own funds requirements for Basel I floor - SA alternative"): "",
        }
        for (table_id, code, description), expected in cases.items():
            with self.subTest(table_id=table_id, code=code, description=description):
                self.assertEqual(builder.dimension_display_format(
                    f"{table_id}.a", "y_axis_rc_code", code, description
                ), expected)

    def test_liquidity_coverage_formats_follow_measure_and_c76_ratio_row(self):
        cases = {
            ("C_72.00", "x_axis_rc_code", "30", "Applicable weight"): "%",
            ("C_72.00", "y_axis_rc_code", "340", "High quality covered bonds (RW35%)"): "",
            ("C_73.00", "x_axis_rc_code", "50", "Weight/Applicable weight"): "%",
            ("C_73.00", "x_axis_rc_code", "60", "Outflow"): "",
            ("C_74.00", "x_axis_rc_code", "10", "Amount/Subject to the 75% cap on inflows"): "",
            ("C_74.00", "x_axis_rc_code", "80", "Applicable weight/Subject to the 75% cap on inflows"): "%",
            ("C_74.00", "x_axis_rc_code", "100", "Applicable weight/Exempted from the cap on inflows"): "%",
            ("C_74.00", "x_axis_rc_code", "140", "Inflow/Subject to the 75% cap on inflows"): "",
            ("C_75.01", "x_axis_rc_code", "50", "Standard weight"): "%",
            ("C_75.01", "x_axis_rc_code", "60", "Applicable weight"): "%",
            ("C_75.01", "x_axis_rc_code", "80", "Inflows subject to the 75% cap on inflows"): "",
            ("C_76.00", "x_axis_rc_code", "10", "Value / Percentage"): "",
            ("C_76.00", "y_axis_rc_code", "30", "Liquidity coverage ratio (%)"): "%",
            ("C_76.00", "y_axis_rc_code", "320", "Inflows Subject to 90% Cap"): "",
        }
        for (table_id, coordinate, code, description), expected in cases.items():
            with self.subTest(table_id=table_id, coordinate=coordinate, code=code):
                self.assertEqual(builder.dimension_display_format(
                    f"{table_id}.a", coordinate, code, description
                ), expected)

    def test_c80_rsf_scale_follows_x_measure_across_sheet_suffixes(self):
        cases = {
            ("x_axis_rc_code", "10", "Amount/Non-HQLA by maturity/< 6 months"): "",
            ("x_axis_rc_code", "40", "Amount/HQLA"): "",
            ("x_axis_rc_code", "90", "Applicable RSF factor/Non-HQLA by maturity/< 6 months"): "%",
            ("x_axis_rc_code", "120", "Applicable RSF factor/HQLA"): "%",
            ("x_axis_rc_code", "130", "Required stable funding"): "",
            ("y_axis_rc_code", "90", "Level 1 assets eligible for 0% LCR haircut"): "",
            ("y_axis_rc_code", "100", "Unencumbered for a residual maturity of less than six months"): "",
        }
        for table_id in ("C_80.00.a", "C_80.00.b", "C_80.00"):
            for (coordinate, code, description), expected in cases.items():
                with self.subTest(table_id=table_id, coordinate=coordinate, code=code):
                    self.assertEqual(builder.dimension_display_format(table_id, coordinate, code, description), expected)

    def test_open_currency_axis_is_curated_without_losing_native_options(self):
        rows = [
            ("qx1", "All currencies"), ("ALL", "Lek"), ("JPY", "Yen"),
            ("EUR", "Euro"), ("USD", "US Dollar"), ("GBP", "Pound Sterling"),
            ("qx46", "Other Currency (open axis tables)"),
        ] + [(f"BA{chr(65 + i)}", f"Currency {i}") for i in range(21)]
        limited = builder.limit_currency_z_rows(rows)
        self.assertEqual([code for code, _ in limited[:5]], ["qx1", "EUR", "USD", "GBP", "JPY"])
        self.assertEqual(limited[-1][0], "qx46")
        self.assertNotIn("ALL", [code for code, _ in limited])
        self.assertLessEqual(sum(len(code) == 3 for code, _ in limited), 21)
        self.assertEqual(
            [code for code, _ in builder.limit_currency_z_rows(rows[:7])],
            ["qx1", "EUR", "USD", "GBP", "JPY", "ALL", "qx46"],
        )

    def test_taxonomy_label_sentence_casing_preserves_regulatory_acronyms(self):
        cases = {
            "DEBT INSTRUMENTS AT COST OR AT AMORTISED COST": "Debt instruments at cost or at amortised cost",
            "TOTAL SREP CAPITAL REQUIREMENT RATIO (TSCR)": "Total SREP capital requirement ratio (TSCR)",
            "MEMORANDUM ITEMS: CAPITAL RATIOS WITHOUT APPLICATION OF IFRS 9": "Memorandum items: capital ratios without application of IFRS 9",
            "OFF-BALANCE SHEET EXPOSURES": "Off-balance sheet exposures",
            "CET1 CAPITAL RATIO": "CET1 capital ratio",
            "SMES AND SMES": "SMEs and SMEs",
            "IS": "IS",
            "Already mixed case": "Already mixed case",
        }
        for source, expected in cases.items():
            with self.subTest(source=source):
                self.assertEqual(builder.format_taxonomy_label(source), expected)

    def test_display_format_inference_marks_percentages_and_raw_units(self):
        cases = {
            "CET1 capital ratio": "%",
            "Average historical annual default rate (%)": "%",
            "Default rate latest year": "%",
            "Total srep capital requirement ratio/To be made up of CET1 capital": "%",
            "Surplus of CET1 capital/CET1 capital ratio without transitional provisions": "%",
            "OCR and Pillar 2 Guidance (P2G)/To be made up of Tier 1 capital": "%",
            "Internal rating scale - PD assigned to the obligor grade or pool": "%",
            "Number of obligors/Of which: defaulted during the year": "Unit",
            "Number of transactions": "Unit",
            "Exposure-weighted average maturity value (days)": "Unit",
            "Cash balances at central banks": "",
            "Derivatives held for trading/Optional: interest rate derivatives": "",
            "CET1 capital ratio/Surplus(+) or deficit(-)": "",
            "Loan to deposit ratio/Denominator": "",
            "Standard conversion factors/10%": "",
        }
        for description, expected in cases.items():
            with self.subTest(description=description):
                self.assertEqual(builder.infer_display_format(description), expected)

        self.assertEqual(
            builder.production_display_format(
                "C_03.00", "y_axis_rc_code", "190", "OCR and Pillar 2 Guidance (P2G)"
            ),
            "%",
        )

    def test_real_dpm_210_delta_and_2020_applicability_rules(self):
        archive = zipfile.ZipFile(SOURCES / "2.10_layouts.zip")
        covid_file = next(n for n in archive.namelist() if "COVID19" in n and n.endswith(".xlsx"))
        fp_file = next(n for n in archive.namelist() if "-FP " in n and n.endswith(".xlsx"))
        self.assertEqual(builder.release_from_name(fp_file, "2.10"), "2.10.1")
        workbook = load_workbook(io.BytesIO(archive.read(covid_file)), data_only=True)
        self.assertIn("F 90.01", workbook.sheetnames)
        sheet = workbook["F 90.01"]
        columns_row, rows_row = builder.find_marker_rows(sheet)
        self.assertTrue(builder.x_axis_rows(sheet, columns_row, rows_row))
        self.assertTrue(builder.y_axis_rows(sheet, rows_row))
        schedule = builder.read_schedule()
        self.assertEqual(builder.schedule_match("COVID19", "2.10", "F_90.01", schedule)["effective_from"], "2020-06-30")
        self.assertEqual(builder.schedule_match("FP", "2.10.1", "Y_01.01", schedule)["effective_from"], "2020-12-31")
        self.assertEqual(builder.schedule_match("SBP", "2.10", "C_106.00", schedule)["effective_from"], "2020-09-30")
        self.assertEqual(builder.schedule_match("FINREP", "3.2", "F_00.01", schedule)["effective_from"], "2022-12-31")
        self.assertEqual(builder.schedule_match("SBP_CR", "4.2", "C_107.00", schedule)["effective_from"], "2025-12-31")
        self.assertEqual(builder.schedule_match("SBPIMV", "4.2", "C_106.00", schedule)["effective_from"], "2026-02-28")
        self.assertEqual(builder.schedule_match("COREP_OF", "4.2", "C_01.00", schedule)["effective_from"], "2026-06-30")
        workbook.close()
        archive.close()

    def test_real_dpm_42_layout_recovers_axes_and_sheet_dimension(self):
        archive = zipfile.ZipFile(SOURCES / "4.2_layouts.zip")
        corep_file = next(n for n in archive.namelist() if "COREP_OFCOREP" in n)
        workbook = load_workbook(io.BytesIO(archive.read(corep_file)), data_only=True)

        sheet = workbook["C_01.00"]
        columns_row, rows_row = builder.find_marker_rows(sheet)
        xs = builder.x_axis_rows(sheet, columns_row, rows_row)
        ys = builder.y_axis_rows(sheet, rows_row)
        self.assertIn(("10", "Amount", 6), xs)
        self.assertIn("10", [code for code, _description, _row in ys])
        self.assertTrue(any(path.startswith("OWN FUNDS") for code, path, _row in ys if code == "10"))
        with (ROOT / "app" / "assets" / "ITS_all_dimension_mapping.csv").open(encoding="utf-8-sig", newline="") as handle:
            current = list(csv.DictReader(handle, delimiter=";"))
        current_y = next(r for r in current if r["table_id"] == "C_01.00" and r["coordinate"] == "y_axis_rc_code" and r["code"] == "10")
        self.assertEqual(current_y["description"].casefold(), "own funds")
        self.assertNotIn("sheet per", " ".join(str(c.value or "") for row in sheet.iter_rows(min_row=1, max_row=5, max_col=12) for c in row).casefold())

        credit_risk = workbook["C_08.01.a"]
        credit_x = builder.x_axis_rows(credit_risk, *builder.find_marker_rows(credit_risk))
        credit_by_code = {code: description for code, description, _row in credit_x}
        self.assertEqual(
            credit_by_code["50"],
            "Credit risk mitigation (CRM) techniques with substitution effects on the exposure/Unfunded credit protection/(-) Credit derivatives",
        )
        self.assertEqual(
            credit_by_code["102"],
            "BREAKDOWN OF THE FULLY ADJUSTED EXPOSURE VALUE OF OFF-BALANCE SHEET ITEMS BY CONVERSION FACTORS/STANDARD CONVERSION FACTORS/0",
        )
        self.assertEqual(
            credit_by_code["103"],
            "BREAKDOWN OF THE FULLY ADJUSTED EXPOSURE VALUE OF OFF-BALANCE SHEET ITEMS BY CONVERSION FACTORS/STANDARD CONVERSION FACTORS/0.1",
        )
        self.assertEqual(
            credit_by_code["90"],
            "Exposure after CRM substitution effects pre conversion factors",
        )
        self.assertEqual(credit_by_code["250"], "Exposure-weighted average maturity value (days)")
        self.assertEqual(credit_by_code["255"], "Risk weighted exposure amount pre supporting factors")
        self.assertEqual(credit_by_code["280"], "Memorandum items:/Expected loss amount")

        operational_risk = workbook["C_17.01.a"]
        operational_x = builder.x_axis_rows(operational_risk, *builder.find_marker_rows(operational_risk))
        self.assertIn(("10", "Event types/INTERNAL FRAUD", 7), operational_x)

        tab_sheet = workbook["C_07.00.a"]
        top = " ".join(str(c.value or "") for row in tab_sheet.iter_rows(min_row=1, max_row=5, max_col=12) for c in row).casefold()
        self.assertIn("sheet per exposure class", top)
        z = builder.dictionary_z_rows(SOURCES / "4.2_dictionary.xlsx", tab_sheet)
        self.assertEqual(len(z), 29)
        z_42_memorandum = builder.dictionary_z_rows(SOURCES / "4.2_dictionary.xlsx", workbook["C_07.00.c"])
        self.assertEqual(len(z_42_memorandum), 8)
        self.assertEqual(len({code for code, _label in z} | {code for code, _label in z_42_memorandum}), 30)
        workbook.close()
        archive.close()

        archive_40 = zipfile.ZipFile(SOURCES / "4.0_layouts.zip")
        corep_file_40 = next(n for n in archive_40.namelist() if "COREP_OFCOREP" in n)
        workbook_40 = load_workbook(io.BytesIO(archive_40.read(corep_file_40)), data_only=True)
        tab_sheet_40 = workbook_40["C_07.00.a"]
        z_40 = builder.dictionary_z_rows(SOURCES / "4.0_dictionary.xlsx", tab_sheet_40)
        self.assertEqual(len(z_40), 29)
        z_40_memorandum = builder.dictionary_z_rows(SOURCES / "4.0_dictionary.xlsx", workbook_40["C_07.00.c"])
        self.assertEqual(len(z_40_memorandum), 8)
        self.assertEqual(len({code for code, _label in z_40} | {code for code, _label in z_40_memorandum}), 30)
        workbook_40.close()
        archive_40.close()

    def test_y_axis_uses_absolute_indent_and_terminal_parent_retroactively(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 2, "Rows")
        rows = [
            ("Cash balances", "5", 2),
            ("Debt securities", "10", 2),
            ("Central banks", "20", 4),
            ("Loans and advances", "70", 2),
            ("Central banks", "80", 4),
            ("Loan commitments", "100", 2),
            ("Households", "110", 4),
            ("Total assets", "120", 0),
        ]
        for row_no, (label, code, indent) in enumerate(rows, start=2):
            sheet.cell(row_no, 2, label).alignment = Alignment(indent=indent)
            sheet.cell(row_no, 3, code)

        rows = builder.y_axis_rows(sheet, 1)
        result = {code: description for code, description, _row in rows}

        self.assertEqual(result["5"], "Total assets/Cash balances")
        self.assertEqual(result["10"], "Total assets/Debt securities")
        self.assertEqual(result["20"], "Total assets/Debt securities/Central banks")
        self.assertEqual(result["70"], "Total assets/Loans and advances")
        self.assertEqual(result["100"], "Total assets/Loan commitments")
        self.assertEqual(result["110"], "Total assets/Loan commitments/Households")
        self.assertEqual(result["120"], "Total assets")
        self.assertLess([code for code, _description, _row in rows].index("120"), [code for code, _description, _row in rows].index("5"))
        workbook.close()

    def test_y_axis_code_less_labels_become_structural_parents(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 2, "Rows")
        for row_no, label, code, indent in [
            (2, "Deferred tax assets and liabilities", "", 0),
            (3, "Total deferred tax assets", "10", 2),
            (4, "Assets not relying on future profitability", "20", 4),
            (5, "Provisions and expected losses", "", 0),
            (6, "IRB excess or shortfall", "30", 2),
        ]:
            sheet.cell(row_no, 2, label).alignment = Alignment(indent=indent)
            sheet.cell(row_no, 3, code)

        entries = builder.y_axis_rows(sheet, 1, include_parent=True)
        by_code = {entry[0]: entry for entry in entries}
        structural_parents = {entry[0]: entry[1] for entry in entries if entry[4]}

        first_parent = by_code["10"][3]
        second_parent = by_code["30"][3]
        self.assertEqual(structural_parents[first_parent], "Deferred tax assets and liabilities")
        self.assertEqual(structural_parents[second_parent], "Provisions and expected losses")
        self.assertEqual(by_code["10"][5], "Deferred tax assets and liabilities/Total deferred tax assets")
        self.assertEqual(by_code["20"][3], "10")
        self.assertEqual(by_code["30"][5], "Provisions and expected losses/IRB excess or shortfall")
        workbook.close()

    def test_real_corep_c03_c04_code_less_section_rows_are_in_hierarchy(self):
        archive = zipfile.ZipFile(SOURCES / "3.2_layouts.zip")
        member = next(name for name in archive.namelist() if "320-P1-COREP 3.2.xlsx" in name)
        workbook = load_workbook(io.BytesIO(archive.read(member)), data_only=True)

        c04_sheet = workbook["C 04.00"]
        c04 = builder.y_axis_rows(c04_sheet, builder.find_marker_rows(c04_sheet)[1], "C_04.00", include_parent=True)
        c04_by_code = {entry[0]: entry for entry in c04}
        self.assertEqual(
            c04_by_code["10"][5],
            "0009 Deferred tax assets and liabilities/Total deferred tax assets",
        )
        self.assertEqual(c04_by_code["20"][3], "10")

        c03_sheet = workbook["C 03.00"]
        c03 = builder.y_axis_rows(c03_sheet, builder.find_marker_rows(c03_sheet)[1], "C_03.00", include_parent=True)
        c03_by_code = {entry[0]: entry for entry in c03}
        self.assertEqual(
            c03_by_code["300"][5],
            "0299 Memorandum Items: Capital ratios without application of the transitional provisions on IFRS 9/CET1 Capital ratio without application of the transitional provisions on IFRS 9",
        )
        self.assertTrue(c03_by_code["300"][3].startswith("__PARENT__"))
        workbook.close()
        archive.close()

    def test_finrep_32_f18_and_f01_hierarchy_from_source_workbook(self):
        archive = zipfile.ZipFile(SOURCES / "3.2_layouts.zip")
        member = next(name for name in archive.namelist() if name.endswith("Annotated Table Layout 321-P2-FINREP 3.2.1.xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(member)), data_only=True)

        f18_rows = builder.y_axis_rows(workbook["F 18.00.a"], 9, "F_18.00.a")
        f18 = {code: description for code, description, _row in f18_rows}
        f18_order = [code for code, _description, _row in f18_rows]
        cost_parent = "DEBT INSTRUMENTS AT COST OR AT AMORTISED COST"
        self.assertEqual(f18["5"], f"{cost_parent}/Cash balances at central banks and other demand deposits")
        self.assertEqual(f18["10"], f"{cost_parent}/Debt securities")
        self.assertEqual(f18["20"], f"{cost_parent}/Debt securities/Central banks")
        self.assertEqual(f18["70"], f"{cost_parent}/Loans and advances")
        self.assertEqual(f18["180"], cost_parent)
        self.assertLess(f18_order.index("180"), f18_order.index("5"))
        fair_value_parent = "DEBT INSTRUMENTS AT FAIR VALUE THROUGH OTHER COMPREHENSIVE INCOME OR THROUGH EQUITY SUBJECT TO IMPAIRMENT"
        self.assertEqual(f18["201"], fair_value_parent)
        self.assertEqual(f18["211"], f"{fair_value_parent}/Debt securities")
        self.assertLess(f18_order.index("201"), f18_order.index("211"))

        f18_off_balance = {code: description for code, description, _row in builder.y_axis_rows(workbook["F 18.00.b"], 9, "F_18.00.b")}
        self.assertEqual(f18_off_balance["340"], "OFF-BALANCE SHEET EXPOSURES/Loan commitments given")
        self.assertEqual(f18_off_balance["350"], "OFF-BALANCE SHEET EXPOSURES/Loan commitments given/Central banks")
        self.assertEqual(f18_off_balance["550"], "OFF-BALANCE SHEET EXPOSURES")

        f01_rows = builder.y_axis_rows(workbook["F 01.01"], 7)
        f01 = {code: description for code, description, _row in f01_rows}
        f01_order = [code for code, _description, _row in f01_rows]
        self.assertEqual(f01["10"], "Total assets/Cash, cash balances at central banks and other demand deposits")
        self.assertEqual(f01["50"], "Total assets/Financial assets held for trading")
        self.assertEqual(f01["60"], "Total assets/Financial assets held for trading/Derivatives")
        self.assertEqual(f01["380"], "Total assets")
        self.assertLess(f01_order.index("380"), f01_order.index("10"))

        workbook.close()
        archive.close()

    def test_finrep_f18_cash_balance_missing_indent_is_corrected_in_42_releases(self):
        archive = zipfile.ZipFile(SOURCES / "4.2_layouts.zip")
        members = (
            next(n for n in archive.namelist() if "FINREP9FINREP 4.2.xlsx" in n),
            next(n for n in archive.namelist() if "FINREP9DPFINREP 4.2.1.xlsx" in n),
        )

        for member in members:
            with self.subTest(workbook=Path(member).name):
                workbook = load_workbook(io.BytesIO(archive.read(member)), data_only=True)
                sheet = workbook["F_18.00.a"]
                rows_row = builder.find_marker_rows(sheet)[1]
                result_rows = builder.y_axis_rows(sheet, rows_row, "F_18.00.a")
                result = {code: description for code, description, _row in result_rows}
                order = [code for code, _description, _row in result_rows]
                parent = "DEBT INSTRUMENTS AT COST OR AT AMORTISED COST"
                self.assertEqual(result["5"], f"{parent}/Cash balances at central banks and other demand deposits")
                self.assertEqual(result["10"], f"{parent}/Debt securities")
                self.assertEqual(result["70"], f"{parent}/Loans and advances")
                self.assertEqual(result["180"], parent)
                self.assertLess(order.index("180"), order.index("5"))
                workbook.close()

        archive.close()

    def test_f18_indent_exception_does_not_apply_to_other_templates(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 2, "Rows")
        sheet.cell(2, 2, "Cash balances").alignment = Alignment(indent=0)
        sheet.cell(2, 3, "5")
        sheet.cell(3, 2, "Debt securities").alignment = Alignment(indent=2)
        sheet.cell(3, 3, "10")

        result = {code: description for code, description, _row in builder.y_axis_rows(sheet, 1, "X_18.00")}

        self.assertEqual(result["5"], "Cash balances")
        self.assertEqual(result["10"], "Cash balances/Debt securities")
        workbook.close()

    def test_parent_codes_follow_structure_even_when_labels_contain_slashes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 2, "Rows")
        sheet.cell(2, 2, "Assets").alignment = Alignment(indent=0)
        sheet.cell(2, 3, "10")
        sheet.cell(3, 2, "Loans / advances").alignment = Alignment(indent=2)
        sheet.cell(3, 3, "20")
        sheet.cell(4, 2, "Retail / SME").alignment = Alignment(indent=4)
        sheet.cell(4, 3, "30")

        rows = builder.y_axis_rows(sheet, 1, include_parent=True)
        by_code = {entry[0]: entry for entry in rows}

        self.assertEqual(by_code["10"][3], "")
        self.assertEqual(by_code["20"][3], "10")
        self.assertEqual(by_code["30"][3], "20")
        self.assertEqual(by_code["30"][1], "Retail / SME")
        self.assertEqual(by_code["30"][5], "Assets/Loans / advances/Retail / SME")
        workbook.close()

    def test_x_parent_code_uses_header_levels(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 1, "Columns")
        sheet.cell(2, 2, "Exposure")
        sheet.cell(3, 2, "Gross amount")
        sheet.cell(3, 3, "Of which")
        sheet.cell(4, 4, "Retail / SME")
        sheet.cell(5, 2, "10")
        sheet.cell(5, 3, "20")
        sheet.cell(5, 4, "30")
        sheet.cell(6, 2, "Rows")
        sheet.merge_cells("B2:D2")
        sheet.merge_cells("C3:D3")

        rows = builder.x_axis_rows(sheet, 1, 6, include_parent=True)
        by_code = {entry[0]: entry for entry in rows}

        self.assertTrue(by_code["20"][3].startswith("__PARENT__"))
        self.assertEqual(by_code["30"][3], "20")
        self.assertTrue(by_code["30"][1].endswith("Retail / SME"))
        virtual_parent = by_code[by_code["20"][3]]
        self.assertEqual(virtual_parent[1], "Exposure")
        self.assertTrue(virtual_parent[4])
        workbook.close()

        # Merge dimensions are presentation only. The same level anchors must
        # reconstruct the same paths when the cells are left unmerged.
        unmerged = Workbook()
        sheet = unmerged.active
        sheet.cell(1, 1, "Columns")
        sheet.cell(2, 2, "Exposure")
        sheet.cell(3, 2, "Gross amount")
        sheet.cell(3, 3, "Of which")
        sheet.cell(4, 4, "Retail / SME")
        sheet.cell(5, 2, "10")
        sheet.cell(5, 3, "20")
        sheet.cell(5, 4, "30")
        sheet.cell(6, 2, "Rows")
        unmerged_rows = builder.x_axis_rows(sheet, 1, 6, include_parent=True)
        self.assertEqual(
            {entry[0]: (entry[1], entry[3], entry[4], entry[5]) for entry in rows},
            {entry[0]: (entry[1], entry[3], entry[4], entry[5]) for entry in unmerged_rows},
        )
        unmerged.close()

    def test_corep_40_unmerged_headers_infer_parent_scopes(self):
        archive = zipfile.ZipFile(SOURCES / "4.0_layouts.zip")
        corep_file = next(n for n in archive.namelist() if "COREP_OFCOREP" in n and n.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(corep_file)), data_only=True)
        sheet = workbook["C_08.01.a"]

        x_rows = builder.x_axis_rows(sheet, *builder.find_marker_rows(sheet))
        x_by_code = {code: description for code, description, _row in x_rows}
        self.assertEqual(
            x_by_code["30"],
            "Original exposure pre conversion factors/Of which: large financial sector entities and unregulated financial entities",
        )
        self.assertEqual(
            x_by_code["50"],
            "Credit risk mitigation (CRM) techniques with substitution effects on the exposure/Unfunded credit protection/(-) Credit derivatives",
        )
        self.assertEqual(
            x_by_code["102"],
            "BREAKDOWN OF THE FULLY ADJUSTED EXPOSURE VALUE OF OFF-BALANCE SHEET ITEMS BY CONVERSION FACTORS/STANDARD CONVERSION FACTORS/0",
        )
        self.assertEqual(
            x_by_code["103"],
            "BREAKDOWN OF THE FULLY ADJUSTED EXPOSURE VALUE OF OFF-BALANCE SHEET ITEMS BY CONVERSION FACTORS/STANDARD CONVERSION FACTORS/0.1",
        )
        self.assertEqual(x_by_code["90"], "Exposure after CRM substitution effects pre conversion factors")

        workbook.close()
        archive.close()

    def test_legacy_corep_0801_z_axis_uses_numbered_template_sheets(self):
        archive = zipfile.ZipFile(SOURCES / "3.2_layouts.zip")
        corep_file = next(n for n in archive.namelist() if "COREP" in n and n.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(corep_file)), data_only=True)
        sheets = [sheet for sheet in workbook.worksheets if sheet.title.startswith("C 08.01.a(")]

        z_rows = [entry for sheet in sheets for entry in builder.legacy_sheet_z_rows(sheet)]
        self.assertEqual(len(sheets), 17)
        self.assertEqual(len(z_rows), 17)
        self.assertEqual(len({code for code, _description in z_rows}), 17)
        self.assertTrue(all(code.isdigit() for code, _description in z_rows))
        self.assertEqual(z_rows[0][0], "0001")
        self.assertIn("Total with own estimates", z_rows[0][1])
        self.assertFalse(any(code.casefold().startswith("qx") for code, _description in z_rows))

        x_rows = builder.x_axis_rows(sheets[0], *builder.find_marker_rows(sheets[0]))
        x_by_code = {code: description for code, description, _row in x_rows}
        self.assertEqual(
            x_by_code["50"],
            "Credit risk mitigation (CRM) techniques with substitution effects on the exposure/Unfunded credit protection/(-) Credit derivatives",
        )

        # The old domain-wide resolver leaked more than 300 unrelated members
        # into this 17-tab template; DPM 1.0 must use its sheets instead.
        leaked = builder.dictionary_z_rows(SOURCES / "3.2_dictionary.xlsx", sheets[0])
        self.assertGreater(len(leaked), 300)

        workbook.close()
        archive.close()

    def test_suffix_is_kept_only_when_it_disambiguates_an_axis_code(self):
        mappings = [
            {"module_code": "COREP", "framework": "3.2", "table_id": "C_01.00.a", "coordinate": "y_axis_rc_code", "code": "10", "description": "Own funds"},
            {"module_code": "COREP", "framework": "3.2", "table_id": "C_01.00.b", "coordinate": "y_axis_rc_code", "code": "10", "description": " own   funds "},
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.a", "coordinate": "y_axis_rc_code", "code": "511", "description": "Own funds/CET1/Share capital"},
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.b", "coordinate": "y_axis_rc_code", "code": "511", "description": "Liabilities/Residual liabilities/Share capital"},
        ]
        templates = [
            {"module_code": "COREP", "framework": "3.2", "template_id": "C_01.00.a", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
            {"module_code": "COREP", "framework": "3.2", "template_id": "C_01.00.b", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
            {"module_code": "RES", "framework": "3.2", "template_id": "T_01.00.a", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
            {"module_code": "RES", "framework": "3.2", "template_id": "T_01.00.b", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
        ]
        exceptions, conflicts = builder.find_discriminating_suffix_templates(mappings, templates)
        self.assertEqual(exceptions, {("RES", "T_01.00")})
        self.assertEqual(conflicts[("RES", "T_01.00")], 1)

    def test_suffix_is_removed_when_other_axes_make_cells_disjoint(self):
        mappings = [
            {"module_code": "FINREP", "framework": "3.2", "table_id": "F_18.00.a", "coordinate": "x_axis_rc_code", "code": "10", "description": "Gross carrying amount"},
            {"module_code": "FINREP", "framework": "3.2", "table_id": "F_18.00.a", "coordinate": "y_axis_rc_code", "code": "340", "description": "Debt instruments held for sale/Loan commitments given"},
            {"module_code": "FINREP", "framework": "3.2", "table_id": "F_18.00.e", "coordinate": "x_axis_rc_code", "code": "20", "description": "Nominal amount"},
            {"module_code": "FINREP", "framework": "3.2", "table_id": "F_18.00.e", "coordinate": "y_axis_rc_code", "code": "340", "description": "Loan commitments given"},
        ]
        templates = [
            {"module_code": "FINREP", "framework": "3.2", "template_id": f"F_18.00.{suffix}", "effective_from": "2022-12-31", "effective_to": "", "status": "official"}
            for suffix in ("a", "e")
        ]

        exceptions, conflicts = builder.find_discriminating_suffix_templates(mappings, templates)

        self.assertNotIn(("FINREP", "F_18.00"), exceptions)
        self.assertNotIn(("FINREP", "F_18.00"), conflicts)

    def test_suffix_is_kept_when_a_complete_cell_key_collides(self):
        mappings = [
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.a", "coordinate": "x_axis_rc_code", "code": "10", "description": "Amount"},
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.a", "coordinate": "y_axis_rc_code", "code": "511", "description": "Assets/Share capital"},
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.b", "coordinate": "x_axis_rc_code", "code": "10", "description": "Amount"},
            {"module_code": "RES", "framework": "3.2", "table_id": "T_01.00.b", "coordinate": "y_axis_rc_code", "code": "511", "description": "Liabilities/Share capital"},
        ]
        templates = [
            {"module_code": "RES", "framework": "3.2", "template_id": f"T_01.00.{suffix}", "effective_from": "2022-12-31", "effective_to": "", "status": "official"}
            for suffix in ("a", "b")
        ]

        exceptions, conflicts = builder.find_discriminating_suffix_templates(mappings, templates)

        self.assertEqual(exceptions, {("RES", "T_01.00")})
        self.assertEqual(conflicts[("RES", "T_01.00")], 1)

    def test_corep_0801_suffixes_merge_when_codes_do_not_collide(self):
        mappings = [
            {"module_code": "COREP", "framework": "2.9.1.1", "table_id": "C_08.01.a", "coordinate": "y_axis_rc_code", "code": "80", "description": "Total exposures/Specialized lending slotting criteria (b)"},
            {"module_code": "COREP", "framework": "2.9.1.1", "table_id": "C_08.01.b", "coordinate": "y_axis_rc_code", "code": "80", "description": "Total exposures/Specialized lending slotting criteria: total"},
            {"module_code": "COREP", "framework": "2.9.1.1", "table_id": "C_08.01.a", "coordinate": "y_axis_rc_code", "code": "100", "description": "Total exposures/Specialized lending slotting criteria (b)/0.5"},
            {"module_code": "COREP", "framework": "2.9.1.1", "table_id": "C_08.01.b", "coordinate": "y_axis_rc_code", "code": "100", "description": "Total exposures/Specialized lending slotting criteria: total/0.5"},
            {"module_code": "COREP", "framework": "2.9.1.1", "table_id": "C_08.01.b", "coordinate": "x_axis_rc_code", "code": "100", "description": "Of which: off balance sheet items"},
        ]
        templates = [
            {"module_code": "COREP", "framework": "2.9.1.1", "template_id": f"C_08.01.{suffix}", "effective_from": "2019-12-31", "effective_to": "2021-06-30", "status": "baseline"}
            for suffix in ("a", "b")
        ]

        exceptions, conflicts = builder.find_discriminating_suffix_templates(mappings, templates)

        self.assertNotIn(("COREP", "C_08.01"), exceptions)
        self.assertNotIn(("COREP", "C_08.01"), conflicts)

    def test_module_schedule_selects_reference_date_intervals(self):
        rows = [
            {"module_code": "COREP", "template_id": "C_01.00", "framework": "3.0.1", "effective_from": "2021-06-30", "effective_to": "2023-06-30", "effective_source": "eba", "status": "official"},
            {"module_code": "COREP", "template_id": "C_01.00", "framework": "3.2", "effective_from": "2023-06-30", "effective_to": "", "effective_source": "eba", "status": "official"},
        ]
        self.assertEqual(resolve("COREP", "C_01.00", builder.date(2022, 12, 31), rows)[0]["framework"], "3.0.1")
        self.assertEqual(resolve("COREP", "C_01.00", builder.date(2023, 6, 30), rows)[0]["framework"], "3.2")

    def test_interval_closure_does_not_cross_module_boundaries(self):
        rows = [
            {"module_code": "COREP", "template_id": "C_01.00", "framework": "3.2", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
            {"module_code": "COREP_OF", "template_id": "C_01.00", "framework": "4.0", "effective_from": "2025-03-31", "effective_to": "", "status": "official"},
            {"module_code": "COREP_OF", "template_id": "C_01.00", "framework": "4.2", "effective_from": "2026-06-30", "effective_to": "", "status": "official"},
            {"module_code": "CODIS", "template_id": "K_04.00.a", "framework": "4.1", "effective_from": "2025-06-30", "effective_to": "", "status": "official"},
            {"module_code": "PILLAR3", "template_id": "K_04.00.a", "framework": "3.3", "effective_from": "2023-12-31", "effective_to": "", "status": "official"},
        ]
        history = builder.close_template_intervals(rows)
        corep = [r for r in history if r["module_code"] == "COREP"]
        corep_of = [r for r in history if r["module_code"] == "COREP_OF"]
        codis = next(r for r in history if r["module_code"] == "CODIS")
        pillar = next(r for r in history if r["module_code"] == "PILLAR3")
        self.assertEqual(corep[0]["effective_to"], "")
        self.assertEqual(corep_of[0]["effective_to"], "2026-06-30")
        self.assertEqual(codis["effective_to"], "")
        self.assertEqual(pillar["effective_to"], "")
        self.assertEqual(resolve("COREP", "C_01.00", builder.date(2026, 7, 1), history)[0]["framework"], "4.2")

    def test_resolver_follows_dpm2_module_split_for_same_template(self):
        rows = [
            {"module_code": "COREP", "template_id": "C_01.00", "framework": "3.2", "effective_from": "2023-06-30", "effective_to": "2026-03-31", "effective_source": "eba", "status": "official"},
            {"module_code": "COREP_OF", "template_id": "C_01.00", "framework": "4.2", "effective_from": "2026-03-31", "effective_to": "", "effective_source": "eba", "status": "official"},
        ]
        self.assertEqual(resolve("COREP", "C_01.00", builder.date(2026, 4, 30), rows)[0]["framework"], "4.2")

    def test_resolver_keeps_parallel_split_modules_ambiguous(self):
        rows = [
            {"module_code": "COREP", "template_id": "C_00.01", "framework": "3.2", "effective_from": "2023-06-30", "effective_to": "", "status": "official"},
            {"module_code": "COREP_OF", "template_id": "C_00.01", "framework": "4.2", "effective_from": "2026-06-30", "effective_to": "", "status": "official"},
            {"module_code": "COREP_LR", "template_id": "C_00.01", "framework": "4.2", "effective_from": "2026-03-31", "effective_to": "", "status": "official"},
        ]
        found = resolve("COREP", "C_00.01", builder.date(2026, 7, 1), rows)
        self.assertEqual({r["module_code"] for r in found}, {"COREP_OF", "COREP_LR"})

    def test_sheet_level_exception_rules_are_explicit(self):
        rows = builder.read_schedule()
        imv = builder.schedule_match("SBP", "3.2.1", "C_106.00", rows)
        other = builder.schedule_match("SBP", "3.2.1", "C_107.01.a", rows)
        self.assertEqual(imv["effective_from"], "2022-09-30")
        self.assertEqual(other["effective_from"], "2022-12-31")
        pay_gap_ci = builder.schedule_match("REM", "3.2.2", "R_06.00.a", rows)
        pay_gap_if = builder.schedule_match("REM", "3.2.2", "R_06.01.a", rows)
        regular_rem = builder.schedule_match("REM", "3.2.2", "R_05.00", rows)
        self.assertEqual(pay_gap_ci["effective_from"], "2023-12-31")
        self.assertEqual(pay_gap_if["effective_from"], "2022-12-31")
        self.assertEqual(regular_rem["effective_from"], "2022-12-31")
        dora = builder.schedule_match("DORA", "4.2", "D_01.00", rows)
        self.assertEqual(dora["status"], "not_applicable_replaced")


if __name__ == "__main__":
    unittest.main()
